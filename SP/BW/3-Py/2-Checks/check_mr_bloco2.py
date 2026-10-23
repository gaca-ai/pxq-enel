# -*- coding: utf-8 -*-
"""
Check PxQ bloco 2 x Mercado Receita (MR), by market option and tariff subgroup.

Question answered: is the PxQ adjustment strategy (bloco 2) aligned with what the
MR shows? One sheet per measure, same rows everywhere (OPCAO x SUBGRUPO, with a
subtotal per OPCAO and a grand total), columns: PxQ [2] | MR | Delta | Delta % |
Participacao no delta total. Deltas, percentages and totals are Excel formulas. ERC (Encargo de Energia de
Reserva) is left out: it is not billed through the BW (known difference).

Sheets:
  0_Resumo            every measure and family, grand totals only, plus the MR
                      reconciliation lines (rounding, subsidy not distributed)
  1_Energia_TUSD_MWh  2_Energia_TE_MWh  3_Demanda_kW  4_Monetario_Total_BRL
  5..11               one sheet per family (FIO A, FIO B, ENCARGOS, PERDAS,
                      ENERGIA, RI, OUTROS)

Rules (agreed with the user on 25/09/2026, README.md section 6 item 7):
  1. Captive energy: the BW books the captive consumption on the TE side only,
     so for CATIVO the TUSD energy = the TE energy (ajuda.txt: both fields carry
     the same value). Free market: all energy on the TUSD side, TE = 0.
  2. Quantities are the billed ones (identical in blocos 1 and 2). MR side uses
     the gross quantities TUSD_E, TE_E and D - not the adjusted Eaj / Daj.
  3. Families: the MR keeps its subsidies (SubsidioTotal) outside the component
     columns, while bloco 2 carries the discounts inside the families. The MR
     computes the subsidy by abating quantity (D - Daj, E - Eaj) times the tariff,
     so the subsidy is spread over the families by the MR's own mechanism: on each
     row, TUSD components x Daj/D (demand rows) or TUSD_Eaj/TUSD_E (energy rows),
     TE components x TE_Eaj/TE_E. No MR row mixes demand and energy with abatement,
     so the split is exact. What the abatement does not explain (tariff-difference
     subsidies of the legacy Baixa Renda tiers) is shown in 0_Resumo.

Same reclassifications as check_mercado_receita.py: LIVRE APE -> LIVRE; Hibrido ->
LIVRE; SD generators billed on a load tariff (EMAE pumping stations) -> LIVRE.

Usage:
    python check_mr_bloco2.py 2026-03
Reads  1-Entrega/pxq/pxq_<ref>.xlsx and the MR of the month (le_mr.py: single-month file or tariff-period file)
Writes 2-Checks/resultados/check_mr_bloco2_<ref>.xlsx
"""
import os
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
sys.path.insert(0, str(Path(__file__).resolve().parent))
from le_mr import FAMILIA_MR, PARCELA_TE, carrega_mr  # noqa: E402
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(os.environ.get("PXQ_BASE", r"C:\python_projects\claude\p_x_q\SP\BW"))
REF = sys.argv[1] if len(sys.argv) > 1 else "2026-03"

ARQ_PXQ = BASE / "1-Entrega" / "pxq" / f"pxq_{REF}.xlsx"
SAIDA = BASE / "2-Checks" / "resultados" / f"check_mr_bloco2_{REF}.xlsx"

NSA = "Não se aplica"
GRAO = ["OPCAO", "SUBGRUPO"]
GRAO_FINO = ["OPCAO", "SUBGRUPO", "MODALIDADE", "POSTO"]
FAMILIAS = ["FIO A", "FIO B", "ENCARGOS", "PERDAS", "ENERGIA", "RI", "OUTROS"]
ORDEM_OPCAO = ["CATIVO", "DISTRIBUIÇÃO", "GERAÇÃO", "LIVRE", "SUPRIMENTO"]
ORDEM_SUBGRUPO = ["A2", "A3a", "A4", "AS", "B1", "B2", "B3", "B4"]

# ------------------------------------------------------------------ MR side
# MR rows of the month: single-month file or tariff-period file (le_mr.py)
mr, ARQ_MR = carrega_mr(BASE, REF)
print("Mercado Receita:", ARQ_MR.name)
mr = mr[mr["Subgrupo"].notna()].copy()
mr["OPCAO"] = mr["OPÇÃO"].replace({"LIVRE APE": "LIVRE"})
# ERC (Encargo de Energia de Reserva) is not billed through the BW - known and
# accepted difference (README.md 6.5), left out of this check
mr = mr[mr["OPCAO"] != "ERC"].copy()
mr["SUBGRUPO"] = mr["Subgrupo"]
for c in ["D", "Daj", "TUSD_E", "TUSD_Eaj", "TE_E", "TE_Eaj", "FinTotAj",
          "Total com Tar de aplicação", "SubsidioTotal"]:
    mr[c] = mr[c].fillna(0.0)

faltam = [c for lista in FAMILIA_MR.values() for c in lista if c not in mr.columns]
assert not faltam, f"colunas do MR nao encontradas: {faltam}"

# rule 3: the subsidy is quantity abated x tariff; no row mixes demand and energy
# with abatement (checked below), so one TUSD factor per row is exact
abate_d = (mr["D"] - mr["Daj"]).abs() > 1e-6
abate_e = (mr["TUSD_E"] - mr["TUSD_Eaj"]).abs() > 1e-6
misto = abate_d & (mr["TUSD_E"] != 0) | abate_e & (mr["D"] != 0)
assert not misto.any(), f"{int(misto.sum())} linha(s) do MR misturam demanda e energia com abatimento"
fator_tusd = np.where(mr["D"] != 0, mr["Daj"] / mr["D"].replace(0, np.nan),
                      np.where(mr["TUSD_E"] != 0, mr["TUSD_Eaj"] / mr["TUSD_E"].replace(0, np.nan), 1.0))
fator_te = np.where(mr["TE_E"] != 0, mr["TE_Eaj"] / mr["TE_E"].replace(0, np.nan), 1.0)
mr["FATOR_TUSD"] = np.nan_to_num(fator_tusd, nan=1.0)
mr["FATOR_TE"] = np.nan_to_num(fator_te, nan=1.0)

# family results go to "FAM_<family>" columns: plain family names would overwrite
# MR columns with the same name ("OUTROS", "RI")
bruto = mr[[c for l in FAMILIA_MR.values() for c in l]].fillna(0.0).sum(axis=1)
for fam, lista in FAMILIA_MR.items():
    te = [c for c in lista if c in PARCELA_TE]
    tusd = [c for c in lista if c not in PARCELA_TE]
    mr[f"FAM_{fam}"] = (mr[tusd].fillna(0.0).sum(axis=1) * mr["FATOR_TUSD"]
               + mr[te].fillna(0.0).sum(axis=1) * mr["FATOR_TE"])

# reconciliation: families + rounding + subsidy not explained by abatement = FinTotAj
mr["NAO_ALOCADO"] = mr["Total com Tar de aplicação"] - bruto
mr["SUBSIDIO_NAO_DISTRIBUIDO"] = (mr["FinTotAj"] - mr[[f"FAM_{f}" for f in FAMILIAS]].sum(axis=1)
                                  - mr["NAO_ALOCADO"])

mr["ENERGIA_TUSD_MWH"] = mr["TUSD_E"]
mr["ENERGIA_TE_MWH"] = mr["TE_E"]
mr["DEMANDA_KW"] = mr["D"]
mr["MONETARIO_TOTAL"] = mr["FinTotAj"]

# ------------------------------------------------------------------ PxQ side
OPCAO_PXQ = {"Cliente Cativo": "CATIVO", "Cliente Livre": "LIVRE", "Cliente Híbrido": "LIVRE",
             "GERAÇÃO": "GERAÇÃO", "USO DISTRIBUIÇÃO": "DISTRIBUIÇÃO", "SUPRIMENTO": "SUPRIMENTO"}
PADRAO_INICIO = re.compile(r"^(A2|A3a|A4|AS|B1|B2|B3|B4)"
                           r"(Azul|Verde|Branca|Convencional|Geração|Distribuição)")
POSTOS = [("EINT", "Intermediário"), ("DINT", "Intermediário"), ("EFP", "Fora ponta"),
          ("DFP", "Fora ponta"), ("EP", "Ponta"), ("DP", "Ponta"), ("E", NSA), ("D", NSA)]


def le_chave(chave):
    """Split a PxQ key into subgrupo, modalidade and posto."""
    m = PADRAO_INICIO.match(chave)
    if not m:
        return pd.Series([None, None, None])
    corpo = chave[:-3] if chave.endswith("MWh") else chave[:-2]
    posto = NSA
    for sufixo, nome in POSTOS:
        if corpo.endswith(sufixo):
            posto = nome
            break
    return pd.Series([m.group(1), m.group(2), posto])


def prepara(df):
    """Tariff-eligible rows with OPCAO / SUBGRUPO / MODALIDADE / POSTO derived."""
    df = df[(df["Tarifável?"] == "SIM") & df["Chave"].notna()].copy()
    df["OPCAO"] = df["Flag Cli.Liv./ Cat."].map(OPCAO_PXQ)
    assert df["OPCAO"].notna().all(), df.loc[df["OPCAO"].isna(), "Flag Cli.Liv./ Cat."].unique()
    df[["SUBGRUPO", "MODALIDADE", "POSTO"]] = df["Chave"].apply(le_chave)
    assert df["SUBGRUPO"].notna().all(), df.loc[df["SUBGRUPO"].isna(), "Chave"].unique()
    carga = (df["OPCAO"] == "GERAÇÃO") & (df["MODALIDADE"] != "Geração")
    df.loc[carga, "OPCAO"] = "LIVRE"
    return df


px = prepara(pd.read_excel(ARQ_PXQ, sheet_name="Componentes"))
fam_px = prepara(pd.read_excel(ARQ_PXQ, sheet_name="Familias"))

# rule 1: energy per fine grain. CATIVO: TUSD = TE = the booked quantity (BW books
# it on one side only); LIVRE: all on TUSD, TE = 0; SD rows keep their own sides.
px["TUSD_MWH"] = px["Consumo TUSD (kWh)"].fillna(0.0) / 1000.0
px["TE_MWH"] = px["Consumo TE (kWh)"].fillna(0.0) / 1000.0
fino = px.groupby(GRAO_FINO, dropna=False)[["TUSD_MWH", "TE_MWH"]].sum().reset_index()
maior = fino[["TUSD_MWH", "TE_MWH"]].max(axis=1)
cativo = fino["OPCAO"] == "CATIVO"
livre = fino["OPCAO"] == "LIVRE"
fino["ENERGIA_TUSD_MWH"] = np.where(cativo | livre, maior, fino["TUSD_MWH"])
fino["ENERGIA_TE_MWH"] = np.where(cativo, maior, np.where(livre, 0.0, fino["TE_MWH"]))

px["DEMANDA_KW"] = px["Demanda (kW)"].fillna(0.0)
px["MONETARIO_TOTAL"] = px["[2] Total TA"].fillna(0.0)
for fam in FAMILIAS:
    fam_px[fam] = fam_px[f"[2] TA · {fam}"].fillna(0.0)

MEDIDAS_PXQ = pd.concat([
    fino.groupby(GRAO)[["ENERGIA_TUSD_MWH", "ENERGIA_TE_MWH"]].sum(),
    px.groupby(GRAO)[["DEMANDA_KW", "MONETARIO_TOTAL"]].sum(),
    fam_px.groupby(GRAO)[FAMILIAS].sum(),
], axis=1).fillna(0.0)
MEDIDAS_MR = mr.groupby(GRAO)[["ENERGIA_TUSD_MWH", "ENERGIA_TE_MWH", "DEMANDA_KW",
                               "MONETARIO_TOTAL"] + [f"FAM_{f}" for f in FAMILIAS]].sum()
MEDIDAS_MR.columns = [c.replace("FAM_", "") for c in MEDIDAS_MR.columns]

# rows: every OPCAO x SUBGRUPO present on either side, in a fixed order
idx = MEDIDAS_PXQ.index.union(MEDIDAS_MR.index)
chaves = sorted(idx, key=lambda k: (ORDEM_OPCAO.index(k[0]) if k[0] in ORDEM_OPCAO else 99,
                                    ORDEM_SUBGRUPO.index(k[1]) if k[1] in ORDEM_SUBGRUPO else 99))
MEDIDAS_PXQ = MEDIDAS_PXQ.reindex(chaves, fill_value=0.0)
MEDIDAS_MR = MEDIDAS_MR.reindex(chaves, fill_value=0.0)

# ------------------------------------------------------------------ workbook
ABAS = [
    ("1_Energia_TUSD_MWh", "ENERGIA_TUSD_MWH", "Energia TUSD (MWh)", "MWh"),
    ("2_Energia_TE_MWh", "ENERGIA_TE_MWH", "Energia TE (MWh)", "MWh"),
    ("3_Demanda_kW", "DEMANDA_KW", "Demanda (kW)", "kW"),
    ("4_Monetario_Total_BRL", "MONETARIO_TOTAL", "Monetário total (R$)", "R$"),
] + [(f"{5 + i}_{f.replace(' ', '_')}", f, f"Família {f} (R$)", "R$") for i, f in enumerate(FAMILIAS)]

AZUL = PatternFill("solid", start_color="1F3864")
CINZA = PatternFill("solid", start_color="D9E1F2")
CINZA2 = PatternFill("solid", start_color="B4C6E7")
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
F = "Arial"
CABEC = ["OPCAO", "SUBGRUPO", "PxQ [2]", "MR", "Delta", "Delta %", "Participação no delta total"]


def formato_num(unidade):
    return '#,##0;[Red]-#,##0;"-"' if unidade != "R$" else '#,##0;[Red]-#,##0;"-"'


def cabecalho(ws, linha, titulos):
    for j, t in enumerate(titulos, 1):
        c = ws.cell(row=linha, column=j, value=t)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = AZUL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA


def aba_medida(wb, nome, medida, titulo, unidade):
    ws = wb.create_sheet(nome)
    ws.cell(row=1, column=1, value=f"{titulo} - PxQ bloco 2 x Mercado Receita - {REF}").font = \
        Font(name=F, bold=True, size=12)
    cabecalho(ws, 3, CABEC)
    linha = 4
    subtotais = []
    total_linha = None
    blocos = []
    for opcao in [o for o in ORDEM_OPCAO if o in {k[0] for k in chaves}] + \
            sorted({k[0] for k in chaves} - set(ORDEM_OPCAO)):
        ini = linha
        for k in [c for c in chaves if c[0] == opcao]:
            ws.cell(row=linha, column=1, value=k[0])
            ws.cell(row=linha, column=2, value=k[1])
            ws.cell(row=linha, column=3, value=float(MEDIDAS_PXQ.loc[k, medida]))
            ws.cell(row=linha, column=4, value=float(MEDIDAS_MR.loc[k, medida]))
            linha += 1
        fim = linha - 1
        ws.cell(row=linha, column=1, value=f"Subtotal {opcao}")
        for col in (3, 4):
            L = get_column_letter(col)
            ws.cell(row=linha, column=col, value=f"=SUM({L}{ini}:{L}{fim})")
        subtotais.append(linha)
        blocos.append((ini, fim, linha))
        linha += 1
    total_linha = linha
    ws.cell(row=total_linha, column=1, value="TOTAL")
    for col in (3, 4):
        L = get_column_letter(col)
        ws.cell(row=total_linha, column=col, value="=" + "+".join(f"{L}{s}" for s in subtotais))
    # deltas, % and share of the total delta - formulas on every row
    for r in range(4, total_linha + 1):
        ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
        ws.cell(row=r, column=6, value=f'=IF(D{r}=0,"",E{r}/D{r})')
        ws.cell(row=r, column=7, value=f'=IF($E${total_linha}=0,"",E{r}/$E${total_linha})')
    # formatting
    for r in range(4, total_linha + 1):
        destaque = CINZA2 if r == total_linha else (CINZA if r in subtotais else None)
        for j in range(1, 8):
            c = ws.cell(row=r, column=j)
            c.font = Font(name=F, bold=destaque is not None, size=10)
            c.border = BORDA
            if destaque:
                c.fill = destaque
            if j in (3, 4, 5):
                c.number_format = formato_num(unidade)
            elif j in (6, 7):
                c.number_format = '0.0%;[Red]-0.0%;"-"'
    for j, w in enumerate([16, 11, 18, 18, 16, 11, 14], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C4"

    return total_linha


wb = Workbook()
wb.remove(wb.active)
resumo = wb.create_sheet("0_Resumo")
totais = {}
for nome, medida, titulo, unidade in ABAS:
    totais[medida] = (nome, aba_medida(wb, nome, medida, titulo, unidade), titulo, unidade)

# summary: grand totals pulled from each sheet by formula
ws = resumo
ws.cell(row=1, column=1, value=f"Resumo - PxQ bloco 2 x Mercado Receita - {REF}").font = \
    Font(name=F, bold=True, size=12)
cabecalho(ws, 3, ["Medida", "Unidade", "PxQ [2]", "MR", "Delta", "Delta %"])
r = 4
linhas_rs = []
for medida, (nome, tl, titulo, unidade) in totais.items():
    ws.cell(row=r, column=1, value=titulo)
    ws.cell(row=r, column=2, value=unidade)
    ws.cell(row=r, column=3, value=f"='{nome}'!C{tl}")
    ws.cell(row=r, column=4, value=f"='{nome}'!D{tl}")
    ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
    ws.cell(row=r, column=6, value=f'=IF(D{r}=0,"",E{r}/D{r})')
    if medida in FAMILIAS:
        linhas_rs.append(r)
    r += 1
# reconciliation of the families against the monetary total
rec = [
    ("Soma das famílias", f"=SUM(C{linhas_rs[0]}:C{linhas_rs[-1]})", f"=SUM(D{linhas_rs[0]}:D{linhas_rs[-1]})"),
    ("MR - resíduo de arredondamento (NAO_ALOCADO)", 0.0, float(mr["NAO_ALOCADO"].sum())),
    ("MR - subsídio não explicado por abatimento de quantidade", 0.0, float(mr["SUBSIDIO_NAO_DISTRIBUIDO"].sum())),
    ("PxQ - resíduo não ajustável (receita sem tarifa no bloco 1)",
     float(px["RECEITA DE FORNECIMENTO"].sum() - px["MONETARIO_TOTAL"].sum()), 0.0),
]
r += 1
ws.cell(row=r, column=1, value="Conciliação das famílias com o monetário total").font = Font(name=F, bold=True)
r += 1
ini_rec = r
for rot, vp, vm in rec:
    ws.cell(row=r, column=1, value=rot)
    ws.cell(row=r, column=2, value="R$")
    ws.cell(row=r, column=3, value=vp)
    ws.cell(row=r, column=4, value=vm)
    ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
    r += 1
ws.cell(row=r, column=1, value="Soma = monetário total (conferência)")
ws.cell(row=r, column=2, value="R$")
lin_mon = 4 + [m for m in totais].index("MONETARIO_TOTAL")
ws.cell(row=r, column=3, value=f"=SUM(C{ini_rec}:C{r - 2})-C{lin_mon}")
ws.cell(row=r, column=4, value=f"=SUM(D{ini_rec}:D{r - 1})-D{lin_mon}")
fim_resumo = r
for rr in range(4, fim_resumo + 1):
    for j in range(1, 7):
        c = ws.cell(row=rr, column=j)
        if c.value is None and j > 2:
            continue
        c.font = Font(name=F, size=10, bold=c.font.bold)
        c.border = BORDA
        if j in (3, 4, 5):
            c.number_format = '#,##0;[Red]-#,##0;"-"'
        elif j == 6:
            c.number_format = '0.0%;[Red]-0.0%;"-"'
for j, w in enumerate([52, 9, 18, 18, 16, 11], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
wb.calculation.fullCalcOnLoad = True
wb.save(SAIDA)

# ------------------------------------------------------------------ console
print(f"check PxQ bloco 2 x Mercado Receita | referencia {REF}")
for nome, medida, titulo, unidade in ABAS:
    a = MEDIDAS_PXQ[medida].sum()
    b = MEDIDAS_MR[medida].sum()
    pct = (a - b) / b if b else float("nan")
    print(f"  {titulo:<26} PxQ [2] {a:>18,.2f} | MR {b:>18,.2f} | delta {a - b:>16,.2f} ({100 * pct:+.2f}%)")
conc_mr = MEDIDAS_MR[FAMILIAS].sum().sum() + mr["NAO_ALOCADO"].sum() + mr["SUBSIDIO_NAO_DISTRIBUIDO"].sum()
print(f"\nconferencia MR: familias + residuos - FinTotAj = {conc_mr - mr['FinTotAj'].sum():,.2f}")
print(f"MR: subsidio nao explicado por abatimento R$ {mr['SUBSIDIO_NAO_DISTRIBUIDO'].sum():,.2f} "
      f"| arredondamento R$ {mr['NAO_ALOCADO'].sum():,.2f}")
print(f"\nsaida: {SAIDA}")
