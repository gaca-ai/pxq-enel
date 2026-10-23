"""
Monta o Mercado Receita (MR) no mesmo layout da Receita Faturada do PxQ e gera o check
celula a celula (PxQ - MR).

Uso:
    python gera_receita_faturada_mr.py 2026-06            # um mes
    python gera_receita_faturada_mr.py 2026-01 2026-08    # varios meses
    python gera_receita_faturada_mr.py --todos            # todos os receita_faturada_<ref>.xlsx
    acrescentar --classes para abrir as linhas por classe (padrao: ABRIR_CLASSES = False)

Por padrao o check e aberto so por opcao de compra x unidade (D/E): as classes sao somadas,
porque o cadastro de classe do livre e diferente no MR e no BW (README 5.4). Com --classes
(ou ABRIR_CLASSES = True) as abas seguem a grade completa do modelo, classe a classe.

Entradas:
- 1-Entrega/receita_faturada/receita_faturada_<ref>.xlsx (lado PxQ, ja no layout)
- src/SP_mercado_receita.xlsx (lado MR, filtrado por AnoMes)
- tbl/tarifas/tarifa_long.csv (tarifas homologadas, so para separar demanda e energia)

Saida: 2-Checks/receita_faturada_MR/receita_faturada_MR_<ref>.xlsx com as abas
PxQ, MR, check (formulas Excel: PxQ - MR, mesma celula) e Conciliacao.
Com mais de um mes, grava tambem consolidado_receita_faturada_MR_<ini>_a_<fim>.xlsx: uma linha
por mes (soma de todas as linhas do layout), nas abas PxQ, MR e check (formulas). A coluna
"Obs. MR" marca o mes em que o MR e igual ao do mes anterior (projetado).

Regras do lado MR:
- Grade: todas as linhas FAT do modelo (opcao x classe x unidade), inclusive as zeradas,
  nas duas abas, para que o check seja uma subtracao celula a celula.
- Opcao: LIVRE APE entra como LIVRE; DISTRIBUICAO = USO D; ERC fica fora (nao passa pelo BW)
  e aparece na Conciliacao. Os clientes GD (Sistema de Compensacao) ficam em CATIVO, como no PxQ.
- Classe: a do MR; Residencial Baixa Renda pela subclasse (baixa renda, Tarifa Social,
  Desconto Social).
- Componentes: de-para da aba Instrucoes do modelo, pelos nomes do MR. CCC, CFURH, LIMINAR_1
  e OUTROS nao tem coluna no modelo e vao para Outros. CVA_* = Base CVA; FinCT_* = Base
  Financeira (FinCT_Geral e CVA_Totaltusd/Totalte sao totais e ficam fora).
- Subsidio: distribuido nos itens pelo abatimento de quantidade do proprio MR (Daj/D na
  demanda, TUSD_Eaj/TUSD_E na energia TUSD, TE_Eaj/TE_E na TE), como no PxQ, que carrega os
  descontos dentro dos itens. O que o abatimento nao explica (diferenca de tarifa da Baixa
  Renda, por exemplo) e distribuido proporcionalmente nos itens da propria linha, que passa
  a somar o FinTotAj do MR (receita depois dos subsidios).
- Demanda x energia: componentes TE vao para E. Componentes TUSD vao para D quando a linha
  so tem demanda e para E quando so tem energia. Nas linhas com as duas (cativo Azul do
  grupo A), cada componente BE e repartido pela tarifa homologada BE do componente
  (demanda x tarifa kW contra energia x tarifa MWh); CVA e FIN pela tarifa BE Total TUSD.
- Tributos: icms em ICMS; cofins + pispasep em PIS/COFINS; repartidos entre D e E na
  proporcao do valor da opcao x classe, como no PxQ.
- Bandeira e UD/ER: o MR nao tem; ficam zerados.
- Geracao na linha E (sem linha no modelo): vai para Outros da linha GERACAO / D, como no PxQ.
- Consumo: MWh na linha E (TE_E no cativo, TUSD_E nas demais opcoes) e MW na D (D / 1000).
- Valores em R$ mil.
"""
import os
import sys
import copy
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

BASE = Path(os.environ.get("PXQ_BASE", r"C:\python_projects\claude\p_x_q\SP\BW"))
DIR_RF = BASE / "1-Entrega" / "receita_faturada"
ARQ_MR = BASE / "src" / "SP_mercado_receita.xlsx"
ARQ_TARIFA = BASE / "tbl" / "tarifas" / "tarifa_long.csv"
DIR_SAIDA = BASE / "2-Checks" / "receita_faturada_MR"
ABRIR_CLASSES = False                                              # True: one row per class (model grid)
ROTULO_SEM_CLASSE = "Todas as classes"

OPCAO_MR = {"CATIVO": "CATIVO", "LIVRE": "LIVRE", "LIVRE APE": "LIVRE",
            "DISTRIBUIÇÃO": "USO D", "GERAÇÃO": "GERAÇÃO", "SUPRIMENTO": "SUPRIMENTO"}
OPCAO_SD = {"USO D", "GERAÇÃO", "SUPRIMENTO"}
CLASSE_MR = {"Comercial": "Comercial", "Industrial": "Industrial", "Poder Público": "Poderes Públicos",
             "Residencial": "Residencial", "Rural": "Rural", "Consumo Próprio": "Consumo Próprio",
             "Serviço Público": "Serviços Públicos", "Iluminação Pública": "Iluminação Pública"}
PADRAO_BAIXA_RENDA = "BAIXA RENDA|TARIFA SOCIAL|DESCONTO SOCIAL"

# MR economic component -> model column (Instrucoes sheet, MR vocabulary)
COLUNA_BE = {
    "DISTRIBUICAO": "Parcela B", "RI": "RI",
    "EnergiaRevenda": "Compra de Energia", "ANGRA": "Compra de Energia",
    "perdastec": "Perdas", "perdasnt": "Perdas", "perdasrb_d": "Perdas", "perdasrb_c": "Perdas",
    "TUST_RB": "Rede Básica", "TUST_FR": "Rede Básica", "TUSDG_T": "Rede Básica",
    "TUSDG_ONS": "Rede Básica", "TUST_ITAIPU": "Rede Básica", "TUST_CI": "Rede Básica",
    "Itaipu": "Transporte Itaipu", "CUSD": "CUSD", "Conexao_D": "CUSD", "Conexao_T": "Conexão T",
    "CDE": "CDE USO", "cDE_U_Covid": "CDE CONTAS", "cDE_E_Covid": "CDE CONTAS", "CDE GD": "CDE CONTAS",
    "Cde_Elet": "CDE ELET", "PROINFA": "PROINFA", "ESSERR": "ESS/ERR/ERCAP", "ONS": "ONS",
    "PeD": "P&D", "TFSEE": "TFSEE",
    "CCC": "Outros", "CFURH": "Outros", "LIMINAR_1": "Outros", "OUTROS": "Outros",
}
# MR components of the TE parcel (always energy)
COMP_TE = {"cDE_E_Covid", "Cde_Elet", "EnergiaRevenda", "ESSERR", "Itaipu", "TUST_CI", "TUST_ITAIPU",
           "perdasrb_c", "CDE GD", "ANGRA", "CFURH"}
# TUSD component -> homologated tariff column (demand / energy split on mixed rows)
TARIFA_COMP = {
    "CDE": "TUSD_CDE", "cDE_U_Covid": "TUSD_CDE_COVID", "Conexao_D": "TUSD_CCD", "Conexao_T": "TUSD_CCT",
    "CUSD": "TUSD_CUSD", "DISTRIBUICAO": "TUSD_FioB", "ONS": "TUSD_ONS", "PROINFA": "TUSD_PROINFA",
    "TFSEE": "TUSD_TFSEE", "TUSDG_ONS": "TUSDG_ONS", "TUSDG_T": "TUSDG_T", "TUST_FR": "TUSD_FR",
    "TUST_RB": "TUSD_RB", "perdasnt": "TUSD_PNT", "perdasrb_d": "TUSD_Per_RB_D", "perdastec": "TUSD_PT",
    "RI": "TUSD_RI", "CCC": "TUSD_CCC", "LIMINAR_1": "Liminar1", "OUTROS": "TUSD Outros", "PeD": "TUSD_PeD",
}
# CVA / FinCT columns of the TE parcel (base name after the prefix, lower case)
BASE_TE = {"cde_e_covid", "energia", "energiarevenda", "esserr", "cfurh", "tranitaipu", "itaipu",
           "tust_ci", "tust_itaipu", "perdasrb_c", "cde elet", "cde gd", "angra"}
TOTAIS_MR = {"CVA_Totaltusd", "CVA_Totalte", "FinCT_Geral"}


def norm(texto):
    """Case- and accent-insensitive key."""
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(t.upper().split())


def num(s):
    return pd.to_numeric(s, errors="coerce").fillna(0.0)


def tarifas_be(ref):
    """BE kW and MWh tariffs of the Azul group-A rows in force in month `ref`."""
    t = pd.read_csv(ARQ_TARIFA, sep=";", decimal=",")
    dia = pd.Timestamp(f"{ref}-15")
    t = t[(pd.to_datetime(t["VIGENCIA_INICIO"]) <= dia) & (pd.to_datetime(t["VIGENCIA_FIM"]) >= dia)]
    t = t[(t["CAMADA"] == "BE") & (t["CLASSE"] == "Não se aplica") & (t["SUBCLASSE"] == "Não se aplica")
          & (t["DETALHE"] == "Não se aplica") & (t["ACESSANTE"] == "Não se aplica")]
    return {(norm(r["SUBGRUPO"]), norm(r["MODALIDADE"]), norm(r["POSTO"]), r["UNIDADE"]): r
            for r in t.to_dict("records")}


def lado_mr(ref, col_valor, grade_d):
    """MR of the month in the model grid: DataFrame indexed by (opcao, classe, unidade), R$."""
    mr = pd.read_excel(ARQ_MR)
    mr = mr[pd.to_datetime(mr["AnoMes"]) == pd.Timestamp(f"{ref}-01")].copy()
    if mr.empty:
        raise ValueError(f"Mercado Receita sem o mes {ref}")
    for c in ["D", "Daj", "TUSD_E", "TUSD_Eaj", "TE_E", "TE_Eaj", "FinTotAj", "icms", "cofins", "pispasep"]:
        mr[c] = num(mr[c])
    info = {"FinTotAj (todas as opcoes)": mr["FinTotAj"].sum()}

    erc = mr["OPÇÃO"] == "ERC"
    info["ERC fora do check (FinTotAj)"] = mr.loc[erc, "FinTotAj"].sum()
    mr = mr[~erc].copy()
    mr["OPCAO"] = mr["OPÇÃO"].map(OPCAO_MR)
    br = mr["Subclasse"].map(norm).str.contains(PADRAO_BAIXA_RENDA)
    mr["CLASSE"] = np.where(mr["OPCAO"].isin(OPCAO_SD), mr["OPCAO"],
                            np.where((mr["Classe"] == "Residencial") & br, "Residencial Baixa Renda",
                                     mr["Classe"].map(CLASSE_MR)))
    sem_classe = mr["OPCAO"].isna() | pd.isna(mr["CLASSE"])
    info["Linhas sem opcao/classe no modelo (FinTotAj)"] = mr.loc[sem_classe, "FinTotAj"].sum()
    mr = mr[~sem_classe].copy()

    # subsidy factors (MR quantity abatement)
    def fator(aj, q):
        return np.where(mr[q] != 0, mr[aj] / mr[q].replace(0, np.nan), 1.0)
    f_d = np.nan_to_num(fator("Daj", "D"), nan=1.0)
    f_eu = np.nan_to_num(fator("TUSD_Eaj", "TUSD_E"), nan=1.0)
    f_te = np.nan_to_num(fator("TE_Eaj", "TE_E"), nan=1.0)

    # demand share of each TUSD component (exact when the row has only D or only E)
    tar = tarifas_be(ref)
    d, e, ete = mr["D"].values, mr["TUSD_E"].values, mr["TE_E"].values
    so_d = (d != 0) & (e == 0)
    misto = (d != 0) & (e != 0)
    chave = list(zip(mr["Subgrupo"].map(norm), mr["Modalidade"].map(norm), mr["Posto"].map(norm)))
    falta_tarifa = 0

    def parte_d(col_tarifa, col_te=None):
        nonlocal falta_tarifa
        s = so_d.astype(float)
        for i in np.where(misto)[0]:
            kw, mwh = tar.get(chave[i] + ("kW",)), tar.get(chave[i] + ("MWh",))
            if kw is None or mwh is None:
                falta_tarifa += 1
                continue
            vd = d[i] * float(kw.get(col_tarifa, 0.0) or 0.0)
            ve = e[i] * float(mwh.get(col_tarifa, 0.0) or 0.0)
            if col_te:
                ve += ete[i] * float(mwh.get(col_te, 0.0) or 0.0)
            s[i] = vd / (vd + ve) if (vd + ve) != 0 else 0.0
        return s

    acum = {}                                                      # (opcao, classe, unidade, coluna) -> R$

    def soma(valores_d, valores_e, coluna):
        for unid, v in (("D", valores_d), ("E", valores_e)):
            g = pd.Series(v, index=mr.index).groupby([mr["OPCAO"], mr["CLASSE"]]).sum()
            for (o, cl), x in g.items():
                if x != 0:
                    acum[(o, cl, unid, coluna)] = acum.get((o, cl, unid, coluna), 0.0) + x

    share_total = parte_d("Total TUSD")
    partes = []                                                    # (value on D, value on E, column), per row
    # economic layer, component by component
    for comp, coluna in COLUNA_BE.items():
        if comp not in mr.columns:
            continue
        v = num(mr[comp]).values
        if comp in COMP_TE:
            vd, ve = np.zeros_like(v), v * f_te
        else:
            s = parte_d(TARIFA_COMP[comp], "TE_PeD" if comp == "PeD" else None)
            vd, ve = v * s * f_d, v * (1 - s) * f_eu
        partes.append((vd, ve, coluna))
    # CVA and FIN layers, totals
    for c in mr.columns:
        if c in TOTAIS_MR or not (c.startswith("CVA_") or c.lower().startswith("finct_")):
            continue
        coluna = "Base CVA" if c.startswith("CVA_") else "Base Financeira"
        base_nome = c.split("_", 1)[1].lower()
        v = num(mr[c]).values
        if base_nome in BASE_TE:
            vd, ve = np.zeros_like(v), v * f_te
        else:
            vd, ve = v * share_total * f_d, v * (1 - share_total) * f_eu
        partes.append((vd, ve, coluna))
    # subsidies the abatement does not explain (e.g. the Baixa Renda tariff difference): each row is
    # scaled to its own FinTotAj, so the MR side carries every discount inside the items
    row_t = sum(vd + ve for vd, ve, _ in partes)
    alvo = mr["FinTotAj"].values
    ok = np.abs(row_t) > 1.0
    k = np.where(ok, alvo / np.where(ok, row_t, 1.0), 1.0)
    info["Subsidio fora do abatimento, distribuido nos itens (R$)"] = float((row_t - alvo)[ok].sum())
    sobra = np.where(ok, 0.0, alvo - row_t)
    info["FinTotAj de linhas sem valor por item, em Outros E (R$)"] = float(sobra.sum())
    for vd, ve, coluna in partes:
        soma(vd * k, ve * k, coluna)
    soma(np.zeros(len(mr)), sobra, "Outros")
    # taxes: split between D and E like the PxQ, by the D/E share of the option x class value
    valor = pd.Series({kk: v for kk, v in acum.items() if kk[3] not in ("CONSUMO",)})
    peso = valor.groupby(level=[0, 1, 2]).sum()
    for coluna, v in (("ICMS", mr["icms"]), ("PIS/COFINS", mr["cofins"] + mr["pispasep"])):
        for (o, cl), x in v.groupby([mr["OPCAO"], mr["CLASSE"]]).sum().items():
            wd, we = peso.get((o, cl, "D"), 0.0), peso.get((o, cl, "E"), 0.0)
            pd_ = wd / (wd + we) if (wd + we) else 0.0
            for unid, p in (("D", pd_), ("E", 1 - pd_)):
                if x * p:
                    acum[(o, cl, unid, coluna)] = acum.get((o, cl, unid, coluna), 0.0) + x * p
    # quantities: MWh on E, MW on D
    q_e = np.where(mr["OPCAO"] == "CATIVO", ete, e)
    soma(d / 1000.0, q_e, "CONSUMO")

    tab = (pd.Series(acum).rename_axis(["OPCAO", "CLASSE", "UNIDADE", "COLUNA"]).unstack("COLUNA")
           .reindex(columns=col_valor).fillna(0.0))
    # generation energy: no E row in the model -> Outros of the D row
    for chave_e in [k for k in tab.index if k[2] == "E" and (k[0], k[1], "D") in grade_d
                    and k not in grade_d]:
        alvo = (chave_e[0], chave_e[1], "D")
        if alvo not in tab.index:
            tab.loc[alvo, :] = 0.0
        tab.loc[alvo, "Outros"] += tab.loc[chave_e].drop("CONSUMO").sum()
        info[f"{' / '.join(chave_e)}: R$ em Outros da linha D (MWh nao gravado)"] = \
            tab.loc[chave_e].drop("CONSUMO").sum()
        tab = tab.drop(index=chave_e)
    fora = [k for k in tab.index if k not in grade_d]
    info["Opcao x classe x unidade fora da grade (R$)"] = tab.loc[fora].drop(columns="CONSUMO").sum().sum()
    info["Linhas com demanda e energia repartidas pela tarifa"] = int(misto.sum())
    info["Linhas mistas sem tarifa encontrada (componente x linha)"] = falta_tarifa
    return tab, info


def le_pxq(ref):
    """PxQ layout of the month: header, (opcao, classe, unidade) -> values, template workbook."""
    wb = openpyxl.load_workbook(DIR_RF / f"receita_faturada_{ref}.xlsx")
    ws = wb["Receita Faturada"]
    cab = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    col = {n: i + 1 for i, n in enumerate(cab) if n}
    dados = {}
    for r in range(4, ws.max_row + 1):
        k = tuple(ws.cell(r, col[n]).value for n in ("OpcaoCompra", "Classe", "Unidade"))
        if all(k):
            dados[k] = {n: ws.cell(r, col[n]).value for n in cab if n}
    return wb, ws, col, dados


def grade_modelo():
    sys.path.insert(0, str(BASE / "3-Py" / "1-Entrega"))
    from gera_receita_faturada import le_modelo                  # noqa: E402
    _, _, _, grade = le_modelo()
    return [(op, cl, u) for _, op, cl, u in grade]


def gera(ref, grade, abrir_classes=ABRIR_CLASSES):
    wb, ws, col, dados = le_pxq(ref)
    id_cols = ["Responsável", "Competência", "TipoFaturamento", "OpcaoCompra", "Classe", "Unidade"]
    nomes = [n for n in col if n not in id_cols and n is not None]
    valor_cols = [n for n in nomes if n != "Total"]              # CONSUMO .. ICMS
    tab_mr, info = lado_mr(ref, valor_cols, set(grade))

    # styles of the first data row, then clear the sheet body
    estilos = [copy.copy(ws.cell(4, c)._style) for c in range(1, ws.max_column + 1)]
    competencia = ws.cell(4, col["Competência"]).value
    ws.delete_rows(4, ws.max_row)
    letra = {n: get_column_letter(c) for n, c in col.items()}
    l_ini, l_fim = letra["Parcela B"], letra["Base CVA"]
    # output rows: model grid, or option x unit with the classes summed
    if abrir_classes:
        linhas_saida = list(grade)
    else:
        linhas_saida = list(dict.fromkeys((op, ROTULO_SEM_CLASSE, u) for op, _, u in grade))
    ultima = 3 + len(linhas_saida)

    def escreve(sheet, origem):
        for i, (op, cl, u) in enumerate(linhas_saida, start=4):
            linha = {"Responsável": "Faturamento", "Competência": competencia, "TipoFaturamento": "FAT",
                     "OpcaoCompra": op, "Classe": cl, "Unidade": u}
            for n in valor_cols:
                linha[n] = origem(op, cl, u, n)
            for n, c in col.items():
                if n in linha:
                    sheet.cell(i, c, linha[n])
            sheet.cell(i, col["Total"], f"=SUM({l_ini}{i}:{l_fim}{i})")
            for c, st in enumerate(estilos, start=1):
                sheet.cell(i, c)._style = copy.copy(st)
            sheet.cell(i, col["Competência"]).number_format = "dd/mm/yyyy"
        t = ultima + 1                                             # totals row
        sheet.cell(t, col["Classe"], "TOTAL")
        for n in valor_cols + ["Total"]:
            L = letra[n]
            sheet.cell(t, col[n], f"=SUM({L}4:{L}{ultima})")
            sheet.cell(t, col[n]).font = openpyxl.styles.Font(bold=True)
            sheet.cell(t, col[n]).number_format = "#,##0.0"
        sheet.cell(t, col["Classe"]).font = openpyxl.styles.Font(bold=True)

    def chaves(op, cl, u, universo):
        """Source keys feeding one output row (one class, or every class of the option)."""
        if cl != ROTULO_SEM_CLASSE:
            return [(op, cl, u)] if (op, cl, u) in universo else []
        return [k for k in universo if k[0] == op and k[2] == u]

    def v_pxq(op, cl, u, n):
        return sum(float(dados[k].get(n) or 0.0) for k in chaves(op, cl, u, dados))

    def v_mr(op, cl, u, n):
        x = sum(float(tab_mr.loc[k, n]) for k in chaves(op, cl, u, set(tab_mr.index)))
        return x if n == "CONSUMO" else x / 1000.0               # R$ -> R$ mil

    ws.title = "PxQ"
    escreve(ws, v_pxq)
    ws_mr = wb.copy_worksheet(ws)
    ws_mr.title = "MR"
    escreve(ws_mr, v_mr)
    ws_ck = wb.copy_worksheet(ws)
    ws_ck.title = "check"
    for i in range(4, ultima + 2):
        for n in valor_cols + ["Total"]:
            L = letra[n]
            ws_ck[f"{L}{i}"] = f"='PxQ'!{L}{i}-'MR'!{L}{i}"
    ws_ck.cell(2, 3, "CHECK = PxQ - MR (R$ mil; consumo em MW/MWh)")
    for nome in [s for s in wb.sheetnames if s not in ("PxQ", "MR", "check")]:
        del wb[nome]

    # reconciliation of the MR side (R$)
    ws_c = wb.create_sheet("Conciliacao")
    gravado_mr = tab_mr.drop(columns="CONSUMO").values.sum()
    linhas = [("Mercado Receita - " + Path(ARQ_MR).name + f" - {ref}", None)]
    linhas += list(info.items())
    linhas += [("Gravado na aba MR (colunas I a AE, R$)", gravado_mr),
               ("Gravado na aba PxQ (colunas I a AE, R$)",
                sum(v_pxq(*k, n) for k in dados for n in valor_cols if n != "CONSUMO") * 1000.0)]
    if ref in ("2026-06",):
        linhas.append(("Atencao: o MR de junho/2026 repete maio (projetado).", None))
    for r in linhas:
        ws_c.append(list(r))
    for row in ws_c.iter_rows():
        for cel in row:
            if isinstance(cel.value, float):
                cel.number_format = "#,##0.00"
    ws_c.column_dimensions["A"].width = 70
    ws_c.column_dimensions["B"].width = 20

    DIR_SAIDA.mkdir(parents=True, exist_ok=True)
    saida = DIR_SAIDA / f"receita_faturada_MR_{ref}.xlsx"
    wb.save(saida)
    print(f"{ref}: MR gravado R$ {gravado_mr/1e6:,.2f} mi | {saida.name}")


def consolida(refs):
    """One row per month (sum of every layout row) for PxQ, MR and check (formulas)."""
    linhas = {"PxQ": [], "MR": []}
    nomes = None
    for ref in refs:
        wb = openpyxl.load_workbook(DIR_SAIDA / f"receita_faturada_MR_{ref}.xlsx", read_only=True)
        for aba in ("PxQ", "MR"):
            rows = list(wb[aba].iter_rows(min_row=3, values_only=True))
            cab = list(rows[0])
            idx = [i for i, n in enumerate(cab)
                   if n and n not in ("Responsável", "Competência", "TipoFaturamento", "OpcaoCompra",
                                      "Classe", "Unidade", "Total")]
            nomes = [cab[i] for i in idx]
            corpo = [r for r in rows[1:] if r[cab.index("OpcaoCompra")]]
            linhas[aba].append([ref] + [sum(float(r[i] or 0.0) for r in corpo
                                             if isinstance(r[i], (int, float))) for i in idx])
    # MR identical to the previous month = projected copy
    obs = [""] + ["projetado (igual ao mês anterior)" if np.allclose(a[1:], b[1:]) else ""
                  for a, b in zip(linhas["MR"][1:], linhas["MR"][:-1])]
    out = openpyxl.Workbook()
    out.remove(out.active)
    cab = ["Mês"] + nomes + ["Total I a Z", "Total geral", "Obs. MR"]
    n_val = len(nomes)
    l_ini = get_column_letter(cab.index("Parcela B") + 1)
    l_z = get_column_letter(cab.index("Base CVA") + 1)
    l_ult = get_column_letter(1 + n_val)
    for aba in ("PxQ", "MR", "check"):
        ws = out.create_sheet(aba)
        ws.append(cab)
        for k, ref in enumerate(refs, start=2):
            ws.cell(k, 1, ref)
            for j in range(n_val):
                c = get_column_letter(j + 2)
                if aba == "check":
                    ws.cell(k, j + 2, f"='PxQ'!{c}{k}-'MR'!{c}{k}")
                else:
                    ws.cell(k, j + 2, linhas[aba][k - 2][j + 1])
                ws.cell(k, j + 2).number_format = "#,##0.0"
            ws.cell(k, n_val + 2, f"=SUM({l_ini}{k}:{l_z}{k})")
            ws.cell(k, n_val + 3, f"=SUM({l_ini}{k}:{l_ult}{k})")
            for j in (n_val + 2, n_val + 3):
                ws.cell(k, j).number_format = "#,##0.0"
            ws.cell(k, n_val + 4, obs[k - 2])
        for c in ws[1]:
            c.font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions[get_column_letter(n_val + 4)].width = 32
        ws.cell(len(refs) + 3, 1, "R$ mil; consumo em MW + MWh (soma das linhas D e E)")
    saida = DIR_SAIDA / f"consolidado_receita_faturada_MR_{refs[0]}_a_{refs[-1]}.xlsx"
    out.save(saida)
    print(f"consolidado: {saida.name}")


if __name__ == "__main__":
    args = sys.argv[1:]
    classes = ABRIR_CLASSES or "--classes" in args
    args = [a for a in args if a != "--classes"]
    if not args:
        sys.exit(__doc__)
    if args == ["--todos"]:
        args = sorted(p.stem[-7:] for p in DIR_RF.glob("receita_faturada_????-??.xlsx"))
    elif len(args) == 2:
        args = [p.strftime("%Y-%m") for p in pd.period_range(args[0], args[1], freq="M")]
    g = grade_modelo()
    for ref in args:
        gera(ref, g, classes)
    if len(args) > 1:
        consolida(args)
