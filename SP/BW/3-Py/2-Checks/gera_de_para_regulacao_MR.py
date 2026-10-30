# -*- coding: utf-8 -*-
"""
Column map PxQ_regulacao (BE/CVA/FIN <item>) -> Mercado Receita (MR) of the
tariff adjustment, item by item and layer by layer, plus a validation by month.

Same rules as check_mr_camadas.py: ERC out of the MR; the MR subsidy is spread
over its components by the MR's own quantity abatement (Daj/D, TUSD_Eaj/TUSD_E,
TE_Eaj/TE_E). Comparison groups put together the PxQ columns that share one MR
column (e.g. TUSD_PeD + TE_PeD -> PeD).

Usage:
    python gera_de_para_regulacao_MR.py 2026-01 2026-05
Reads  1-Entrega/pxq_regulacao/PxQ_regulacao_<ref>.xlsx and the MR of each month (le_mr.py)
Writes 2-Checks/de_para_PxQ_regulacao_MR.xlsx
"""
import os
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from le_mr import FAMILIA_MR, PARCELA_TE, carrega_mr  # noqa: E402

BASE = Path(os.environ.get("PXQ_BASE", r"C:\python_projects\claude\p_x_q\SP\BW"))
INI, FIM = (sys.argv[1], sys.argv[2]) if len(sys.argv) > 2 else ("2026-01", "2026-05")
REFS = [p.strftime("%Y-%m") for p in pd.period_range(INI, FIM, freq="M")]
SAIDA = BASE / "2-Checks" / "de_para_PxQ_regulacao_MR.xlsx"
MR_MIL = 1e6  # R$ -> MR$ (millions of reais)

# PxQ item -> (family, parcel, MR economic column, MR CVA column, MR FinCT column).
# None = the MR has no column for that layer of the item.
ITENS = {
    "TUSD_RB":               ("FIO A",    "TUSD", "TUST_RB",        "CVA_TUST_RB",              "finCT_TUST_RB"),
    "TUSD_FR":               ("FIO A",    "TUSD", "TUST_FR",        "CVA_TUST_FR",              "finCT_TUST_FR"),
    "TUSD_CCT":              ("FIO A",    "TUSD", "Conexao_T",      None,                       "FinCT_Conexao_T"),
    "TUSD_CCD":              ("FIO A",    "TUSD", "Conexao_D",      None,                       "FinCT_Conexao_D"),
    "TUSD_CUSD":             ("FIO A",    "TUSD", "CUSD",           None,                       "FinCT_CUSD"),
    "TUSDG_T":               ("FIO A",    "TUSD", "TUSDG_T",        "CVA_TUSDG_T",              "FINCT_TUSDG_T"),
    "TE_TRANSPORTE_ITAIPU":  ("FIO A",    "TE",   "ITAIPU",         "CVA_TranItaipu",           "FinCT_Itaipu"),
    "TE_TUST_ITAIPU":        ("FIO A",    "TE",   "TUST_ITAIPU",    "CVA_TUST_ITAIPU",          "FINCT_TUST_ITAIPU"),
    "TE_TUST_CI":            ("FIO A",    "TE",   "TUST_CI",        "CVA_TUST_CI",              "FINCT_TUST_CI"),
    "TUSD_FioB":             ("FIO B",    "TUSD", "DISTRIBUICAO",   None,                       "FinCT_ParcelaB"),
    "TUSD_CDE_COVID":        ("ENCARGOS", "TUSD", "covidTusdEco",   "CVA_CovidTusd",            "FinCT_CDECOVIDTUSD"),
    "TUSD_TFSEE":            ("ENCARGOS", "TUSD", "TFSEE",          None,                       "finCT_TFSEE"),
    "TUSD_PeD":              ("ENCARGOS", "TUSD", "PeD",            None,                       "FinCT_Ped"),
    "TE_PeD":                ("ENCARGOS", "TE",   "PeD",            None,                       "FinCT_Ped"),
    "TUSD_ONS":              ("ENCARGOS", "TUSD", "ONS",            None,                       "FinCT_ONS"),
    "TUSD_CCC":              ("ENCARGOS", "TUSD", "CCC",            "CVA_CCC",                  "FinCT_CCC"),
    "TUSD_CDE":              ("ENCARGOS", "TUSD", "CDE",            "CVA_CDE",                  "FinCT_CDE"),
    "TUSD_PROINFA":          ("ENCARGOS", "TUSD", "PROINFA",        "CVA_PROINFA",              "finCT_PROINFA"),
    "Liminar1":              ("ENCARGOS", "TUSD", None,             "CVA_liminar",              "FINCT_Liminar"),
    "TUSDG_ONS":             ("ENCARGOS", "TUSD", "TUSDG_ONS",      "CVA_TUSDG_ONS",            "FINCT_TUSDG_ONS"),
    "TE_CDE_COVID":          ("ENCARGOS", "TE",   "covidTeEco",     "CVA_CovidTe_ou_CVA_CDE_E", "FinCT_CDECOVIDTE_ou_FinCT_CDE_E"),
    "TE_CDE_ELET":           ("ENCARGOS", "TE",   "CDE_Elet_Eco",   "CVA_CDE_Elet",             "FinCT_CDE_Elet"),
    "TE_CDE_GD":             ("ENCARGOS", "TE",   "CDE_GD_Eco",     "CVA_CDE_GD",               "FinCT_CDE_GD"),
    "TE_ESSERR":             ("ENCARGOS", "TE",   "ESSERR",         "CVA_ESSERR",               "FinCT_ESSERR"),
    "TE_CFURH":              ("ENCARGOS", "TE",   "CFURH",          "CVA_CFURH",                "FinCT_CFURH"),
    "TUSD_PT":               ("PERDAS",   "TUSD", "perdastec",      "CVA_perdastec",            "FinCT_perdastec"),
    "TUSD_Per_RB_D":         ("PERDAS",   "TUSD", "perdasrb_d",     "CVA_perdasrb_d",           "FINCT_perdasrb_d"),
    "TUSD_PNT":              ("PERDAS",   "TUSD", "perdasnt",       "CVA_perdasnt",             "FinCT_perdasnt"),
    "TE_Per_RB":             ("PERDAS",   "TE",   "perdasrb_c",     "CVA_perdasrb_c",           "FINCT_perdasrb_c"),
    "TE_ENERGIA":            ("ENERGIA",  "TE",   "EnergiaRevenda", "CVA_EnergiaRevenda",       "FinCT_EnergiaRevenda"),
    "TE_ANGRA":              ("ENERGIA",  "TE",   "ANGRA_Eco",      "CVA_ANGRA",                "FinCT_ANGRA"),
    "TUSD_RI":               ("RI",       "TUSD", "RI",             None,                       "FINCT_RI"),
    "TUSD Outros":           ("OUTROS",   "TUSD", "OUTROS",         None,                       None),
    "TUSD Subsidio":         ("OUTROS",   "TUSD", None,             None,                       "FinCT_Subsidio"),
    "TE Subsidio":           ("OUTROS",   "TE",   None,             None,                       "FinCT_Subsidio"),
    "TUSD BENEFICIO_L14299": ("OUTROS",   "TUSD", "Subs_BENEFICIO_L14299_TUSD", None,           None),
    "TE BENEFICIO_L14299":   ("OUTROS",   "TE",   "Subs_BENEFICIO_L14299_TE",   None,           None),
}
# Notes shown in the map for the non 1:1 cases
NOTAS = {
    "PeD": "Coluna unica TUSD + TE no MR: TUSD_PeD e TE_PeD comparados juntos",
    "FinCT_Ped": "Coluna unica TUSD + TE no MR: TUSD_PeD e TE_PeD comparados juntos",
    "FinCT_Subsidio": "Coluna unica TUSD + TE no MR: TUSD Subsidio e TE Subsidio comparados juntos",
    "CVA_CovidTe_ou_CVA_CDE_E": "Coluna do MR compartilhada com a CDE Energia (CDE_E)",
    "FinCT_CDECOVIDTE_ou_FinCT_CDE_E": "Coluna do MR compartilhada com a CDE Energia (CDE_E)",
    "Liminar1": "Sem base economica no MR (so CVA e FinCT da liminar)",
}
CAMADAS = [("BE", 2), ("CVA", 3), ("FIN", 4)]
FAMILIAS = ["FIO A", "FIO B", "ENCARGOS", "PERDAS", "ENERGIA", "RI", "OUTROS"]


def monta_de_para():
    """Long map: one row per PxQ column x layer, plus MR columns with no PxQ item."""
    linhas = []
    for item, spec in ITENS.items():
        fam, parc = spec[0], spec[1]
        for cam, pos in CAMADAS:
            col_mr = spec[pos]
            linhas.append({
                "Família": fam, "Parcela": parc, "Camada": cam, "Item PxQ": item,
                "Coluna PxQ_regulacao": f"{cam} {item}",
                "Coluna MR": col_mr if col_mr else "(sem coluna no MR)",
                "Fator de subsídio do MR": ("TE_Eaj / TE_E" if col_mr in PARCELA_TE
                                            else "Daj / D ou TUSD_Eaj / TUSD_E") if col_mr else "",
                "Observação": NOTAS.get(col_mr) or NOTAS.get(item, ""),
            })
    usadas = {r["Coluna MR"] for r in linhas}
    for fam, cols in FAMILIA_MR.items():
        for c in cols:
            if c not in usadas:
                cam = "CVA" if c.startswith("CVA_") else "FIN" if c.lower().startswith("finct_") else "BE"
                linhas.append({
                    "Família": fam, "Parcela": "TE" if c in PARCELA_TE else "TUSD", "Camada": cam,
                    "Item PxQ": "(sem item no PxQ)", "Coluna PxQ_regulacao": "(sem coluna no PxQ)",
                    "Coluna MR": c, "Fator de subsídio do MR": "",
                    "Observação": "Coluna do MR sem componente correspondente na tabela de tarifas (PCAT)",
                })
    dp = pd.DataFrame(linhas)
    # comparison group = MR column; PxQ columns without MR column compare alone
    dp["Grupo de comparação"] = np.where(dp["Coluna MR"].str.startswith("("),
                                         dp["Coluna PxQ_regulacao"], dp["Coluna MR"])
    ordem_fam = {f: i for i, f in enumerate(FAMILIAS)}
    ordem_cam = {"BE": 0, "CVA": 1, "FIN": 2}
    dp = dp.sort_values(["Camada", "Família", "Parcela", "Item PxQ"],
                        key=lambda s: s.map(ordem_cam) if s.name == "Camada"
                        else s.map(ordem_fam) if s.name == "Família" else s)
    return dp.reset_index(drop=True)


def mr_por_coluna(ref):
    """MR of the month by component column, subsidy spread by abatement (R$)."""
    mr, _ = carrega_mr(BASE, ref)
    mr = mr[mr["Subgrupo"].notna()].copy()
    mr = mr[mr["OPÇÃO"].replace({"LIVRE APE": "LIVRE"}) != "ERC"].copy()
    for c in ["D", "Daj", "TUSD_E", "TUSD_Eaj", "TE_E", "TE_Eaj"]:
        mr[c] = mr[c].fillna(0.0)
    f_tusd = np.where(mr["D"] != 0, mr["Daj"] / mr["D"].replace(0, np.nan),
                      np.where(mr["TUSD_E"] != 0, mr["TUSD_Eaj"] / mr["TUSD_E"].replace(0, np.nan), 1.0))
    f_te = np.where(mr["TE_E"] != 0, mr["TE_Eaj"] / mr["TE_E"].replace(0, np.nan), 1.0)
    f_tusd, f_te = np.nan_to_num(f_tusd, nan=1.0), np.nan_to_num(f_te, nan=1.0)
    cols = [c for lista in FAMILIA_MR.values() for c in lista]
    return {c: (mr[c].fillna(0.0) * (f_te if c in PARCELA_TE else f_tusd)).sum() for c in cols}


def pxq_por_coluna(ref):
    """PxQ_regulacao of the month by BE/CVA/FIN column (R$)."""
    df = pd.read_excel(BASE / "1-Entrega" / "pxq_regulacao" / f"PxQ_regulacao_{ref}.xlsx")
    cols = [c for c in df.columns if str(c).split(" ", 1)[0] in ("BE", "CVA", "FIN")]
    return df[cols].fillna(0.0).sum().to_dict()


def valida(dp):
    """PxQ x MR by comparison group and month (MR$)."""
    grupos = dp.groupby("Grupo de comparação", sort=False).agg(
        Camada=("Camada", "first"), Família=("Família", "first"),
        cols_pxq=("Coluna PxQ_regulacao", lambda s: [c for c in s if not c.startswith("(")]),
        col_mr=("Coluna MR", lambda s: next((c for c in s if not c.startswith("(")), None)))
    blocos = []
    for ref in REFS:
        pxq, mr = pxq_por_coluna(ref), mr_por_coluna(ref)
        faltam = [c for cs in grupos["cols_pxq"] for c in cs if c not in pxq]
        if faltam:
            raise KeyError(f"{ref}: colunas do de-para ausentes no PxQ_regulacao: {faltam}")
        b = pd.DataFrame({
            "PxQ": [sum(pxq[c] for c in cs) / MR_MIL for cs in grupos["cols_pxq"]],
            "MR": [(mr[c] / MR_MIL) if c else 0.0 for c in grupos["col_mr"]],
        }, index=grupos.index)
        b.columns = pd.MultiIndex.from_product([[ref], b.columns])
        blocos.append(b)
    v = pd.concat(blocos, axis=1)
    pxq_ac = v.xs("PxQ", axis=1, level=1).sum(axis=1)
    mr_ac = v.xs("MR", axis=1, level=1).sum(axis=1)
    v.columns = [f"{r} {lado}" for r, lado in v.columns]
    out = grupos[["Camada", "Família"]].join(v)
    out["Acumulado PxQ"], out["Acumulado MR"] = pxq_ac, mr_ac
    out["Acumulado Δ"] = pxq_ac - mr_ac
    out["Acumulado Δ %"] = (out["Acumulado Δ"] / mr_ac.abs().replace(0, np.nan)).where(mr_ac.abs() > 1e-9)
    out = out.reset_index().rename(columns={"Grupo de comparação": "Grupo (coluna do MR)"})
    tot = out.select_dtypes("number").sum()
    tot["Acumulado Δ %"] = tot["Acumulado Δ"] / abs(tot["Acumulado MR"])
    tot = pd.DataFrame([{"Grupo (coluna do MR)": "TOTAL", "Camada": "", "Família": "", **tot.to_dict()}])
    return pd.concat([out, tot], ignore_index=True)


def por_familia(val):
    """Family view of the validation: same lines as check_mr_camadas.py (MR$)."""
    v = val[val["Grupo (coluna do MR)"] != "TOTAL"].copy()
    v["Linha"] = np.where(v["Camada"] == "BE", "BE · " + v["Família"], v["Camada"] + " (total)")
    linhas = [f"BE · {f}" for f in FAMILIAS] + ["CVA (total)", "FIN (total)"]
    num = [c for c in v.columns if c.endswith((" PxQ", " MR")) or c in ("Acumulado Δ",)]
    f = v.groupby("Linha")[num].sum().reindex(linhas)
    f.loc["TOTAL"] = f.sum()
    f["Acumulado Δ %"] = f["Acumulado Δ"] / f["Acumulado MR"].abs().replace(0, np.nan)
    return f.reset_index()


LEIA_ME = [
    "De-para de colunas: PxQ_regulacao_<ref>.xlsx (BE/CVA/FIN <item>) -> Mercado Receita (MR) do reajuste",
    "",
    "Aba 1_De_para: uma linha por coluna do PxQ_regulacao (37 itens x 3 camadas) e as colunas do MR sem item no PxQ.",
    "Camadas: BE = base economica; CVA = Conta de Compensacao de Variacao de Valores de Itens da Parcela A;",
    "FIN = componentes financeiros (no MR, colunas FinCT_*).",
    "Grupo de comparacao = coluna do MR. Varias colunas do PxQ no mesmo grupo sao somadas (ex.: TUSD_PeD + TE_PeD -> PeD).",
    "Fator de subsidio do MR: cada coluna do MR e multiplicada, linha a linha, pelo abatimento de quantidade do proprio MR,",
    "como em check_mr_camadas.py; assim os descontos ficam dentro dos componentes nos dois lados.",
    "",
    f"Aba 2_Validacao_MR$: PxQ x MR por grupo, {INI} a {FIM}, em MR$ (milhoes de reais). ERC fora do MR.",
    "Delta % = (PxQ - MR) / |MR| no acumulado.",
    "Aba 3_Validacao_familia_MR$: mesma validacao somada por familia (BE) e camada (CVA e FIN totais),",
    "nas mesmas linhas do consolidado de check_mr_camadas.py.",
    "",
    "Gerado por 3-Py/2-Checks/gera_de_para_regulacao_MR.py",
]


def grava(dp, val, fam):
    with pd.ExcelWriter(SAIDA, engine="openpyxl") as w:
        pd.DataFrame({"Leia-me": LEIA_ME}).to_excel(w, sheet_name="Leia-me", index=False)
        dp.to_excel(w, sheet_name="1_De_para", index=False)
        val.to_excel(w, sheet_name="2_Validacao_MR$", index=False)
        fam.to_excel(w, sheet_name="3_Validacao_familia_MR$", index=False)
        ws = w.book["3_Validacao_familia_MR$"]
        ws.column_dimensions["A"].width = 18
        for row in ws.iter_rows(min_row=2, min_col=2):
            for c in row:
                c.number_format = "0.00%" if c.column == ws.max_column else "#,##0.000"
        ws = w.book["Leia-me"]
        ws.column_dimensions["A"].width = 120
        ws["A1"].font = Font(bold=True)
        ws = w.book["1_De_para"]
        for col, larg in zip("ABCDEFGHI", [11, 8, 8, 24, 28, 32, 28, 70, 32]):
            ws.column_dimensions[col].width = larg
        ws = w.book["2_Validacao_MR$"]
        ws.column_dimensions["A"].width = 32
        ultima = ws.max_column
        for row in ws.iter_rows(min_row=2, min_col=4):
            for c in row:
                c.number_format = "0.00%" if c.column == ultima else "#,##0.000"
        for c in ws[1]:
            c.alignment = Alignment(wrap_text=True)
        for nome in w.book.sheetnames:
            for c in w.book[nome][1]:
                c.font = Font(bold=True)


if __name__ == "__main__":
    dp = monta_de_para()
    val = valida(dp)
    fam = por_familia(val)
    grava(dp, val, fam)
    pd.set_option("display.width", 250)
    print(val[["Grupo (coluna do MR)", "Camada", "Acumulado PxQ", "Acumulado MR", "Acumulado Δ", "Acumulado Δ %"]]
          .round(3).to_string(index=False))
    print(fam[["Linha", "Acumulado PxQ", "Acumulado MR", "Acumulado Δ", "Acumulado Δ %"]].round(3).to_string(index=False))
    print(f"\nsaida: {SAIDA}")
