"""
Compara, no layout Receita Faturada, a tarifa homologada com o bloco 2 (entrega) do PxQ e separa a
diferenca por origem. Objetivo: identificar os itens que recebem descontos tarifarios reembolsados
pela CDE (Tarifa Social e fonte incentivada).

Uso:
    python gera_receita_faturada_bl1_vs_bl2.py 2026-06            # um mes
    python gera_receita_faturada_bl1_vs_bl2.py 2026-01 2026-08    # varios meses
    python gera_receita_faturada_bl1_vs_bl2.py --todos            # todos os pxq_<ref>.xlsx

Entradas: 1-Entrega/pxq/pxq_<ref>.xlsx (abas Componentes e Componentes_BE, com os totais por
camada dos blocos 1 e 2 e as colunas de ajuste), 1-Entrega/receita_faturada/receita_faturada_<ref>.xlsx
(so o cabecalho) e src/dmr_tarifa_social.csv (DMR faturada da Tarifa Social, informe de gestao).
Saida: 2-Checks/receita_faturada_Bl1_vs_Bl2/receita_faturada_Bl1_vs_Bl2_<ref>.xlsx, com as abas
  tarifa_homologada    layout pelo bloco 1 (tarifa homologada, com tarifa proporcional; a Tarifa
                       Social pela tarifa B1 subclasse Baixa Renda sem a reducao de 100% ate 80 kWh,
                       PRORET 5.2, par. 129-130) menos as duas parcelas que nao sao reembolsadas pela
                       CDE: a isencao de encargos do APE e o SCEE - Baixa Renda (NOTA_FORA)
  tarifa_realizada     layout pelo bloco 2 (a entrega)
  diferenca            formulas Excel tarifa_homologada - tarifa_realizada, celula a celula
  ta_somada            tarifa de aplicacao somada (colunas I a Z), uma linha por opcao x classe x
                       unidade: tarifa homologada, tarifa realizada, diferenca (formula) e a diferenca
                       por origem
  desconto_tarifa_social_dmr   Tarifa Social - DMR: a DMR faturada do mes, aberta pelos itens na
                               proporcao da etapa Tarifa Social do bloco 2
  desconto_fonte_incentivada   parte que vem do desconto de fonte incentivada (1a passada)
  ajuste_2a_passada            ajuste de fechamento (nao e desconto)
  conciliacao          soma das origens = diferenca, conferencia com as colunas de ajuste do notebook
                       e, como memoria, as duas parcelas retiradas da tarifa homologada
Com mais de um mes, grava tambem consolidado_receita_faturada_Bl1_vs_Bl2_<ini>_a_<fim>.xlsx,
com uma linha por mes em cada aba.

Regras:
- O layout de cada bloco e montado pelo proprio conversor (3-Py/1-Entrega/gera_receita_faturada.py,
  funcao converte), com a mesma grade, de-para e classes da entrega: opcao x classe x unidade,
  todas as linhas do modelo, inclusive as zeradas. Bandeira, UD/ER, Outros e tributos nao dependem
  do bloco e ficam iguais nos dois lados.
- Diferenca positiva = desconto (o bloco 1 e maior que o bloco 2).
- Fora da analise (retiradas da tarifa homologada; valores so na conciliacao, como memoria):
  * Isencao de encargos (APE): PRORET 5.2, par. 81 e 91 - o consumo isento sai do mercado de rateio
    da cota da CDE e nao e reembolsado; o custo fica com os demais consumidores pelo rateio.
  * SCEE - Baixa Renda: e a TE do SCEE, homologada no reajuste, aplicada sobre a energia compensada.
- Separacao por origem, linha a linha do pxq, na ordem das etapas do bloco 2 (Passo 9 do notebook):
  * Tarifa Social: a etapa do notebook (coluna "Ajuste Tarifa Social (R$)"), proporcional aos itens
    do bloco 1 da linha, em duas partes: DMR = a DMR faturada do mes (src/dmr_tarifa_social.csv) e
    SCEE - Baixa Renda = etapa - DMR (fora da analise). Mes sem DMR no arquivo: tudo em DMR (estimado).
    O Desconto Social e a isencao de encargos ficam no bloco 1: nao sao reembolsados pela CDE
    (PRORET 5.2, par. 81 e 91).
  * 2a passada: fator unico sobre todos os itens da linha, f = bloco 2 / (bloco 2 - ajuste da
    2a passada). O valor antes da 2a passada e o bloco 2 / f, item a item e camada a camada.
  * Fonte incentivada: diferenca entre o valor depois da Tarifa Social e o valor antes da 2a
    passada nos itens de Fio A e Fio B (a 1a passada so atua nesses itens).
  * APE (fora da analise): a mesma diferenca nos demais itens (a isencao do APE atua nos encargos).
  * Base CVA e Base Financeira: exatas na Tarifa Social e na 2a passada (mesmo fator). Entre APE e
    fonte incentivada, a soma CVA + FIN de cada origem e exata (sai dos itens); a divisao entre CVA
    e FIN segue a proporcao da linha. So e aproximada quando a mesma linha tem APE e fonte
    incentivada.
- Geradoras, suprimentos e uso da distribuicao (SD) nao passam pelo bloco 2: nao tem origem
  atribuida (os blocos intermediarios sao a propria entrega).
- Valores em R$ mil; consumo em MW (linha D) e MWh (linha E).
"""
import os
import sys
import copy
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

BASE = Path(os.environ.get("PXQ_BASE", r"C:\python_projects\claude\p_x_q\SP\BW"))
sys.path.insert(0, str(BASE / "3-Py" / "1-Entrega"))
import gera_receita_faturada as grf                                # noqa: E402

DIR_PXQ = BASE / "1-Entrega" / "pxq"
DIR_RF = BASE / "1-Entrega" / "receita_faturada"
DIR_SAIDA = BASE / "2-Checks" / "receita_faturada_Bl1_vs_Bl2"

# tariff items of Fio A and Fio B (same families as the notebook, Passo 9)
ITENS_FIO = {"TUSD_RB", "TUSD_FR", "TUSD_CCT", "TUSD_CCD", "TUSD_CUSD", "TUSDG_T",
             "TE_TRANSPORTE_ITAIPU", "TE_TUST_ITAIPU", "TE_TUST_CI", "TUSD_FioB"}
COL_AJ_APE, COL_AJ_1, COL_AJ_2 = "Ajuste regime APE (R$)", "Ajuste 1ª passada (R$)", "Ajuste 2ª passada (R$)"
COL_AJ_TS = "Ajuste Tarifa Social (R$)"
ABA_HOMOL, ABA_REAL, ABA_DIF, ABA_TA, ABA_CONC = ("tarifa_homologada", "tarifa_realizada", "diferenca",
                                                   "ta_somada", "conciliacao")
# (sheet, label, block before, block after)
ABAS_ORIGEM = [("desconto_tarifa_social_dmr", "Tarifa Social – DMR", "[1]", "[D]"),
               ("desconto_fonte_incentivada", "Fonte incentivada", "[A]", "[P]"),
               ("ajuste_2a_passada", "2a passada", "[P]", "[2]")]
# parts of the bloco 1 - bloco 2 gap not reimbursed by the CDE: taken out of the tariff baseline
# ([H] = [1] - SCEE - APE) and shown only as a memo in the reconciliation
FORA = [("SCEE - Baixa Renda", "[D]", "[T]"), ("Isenção de encargos (APE)", "[T]", "[A]")]
ID_COLS = ["Responsável", "Competência", "TipoFaturamento", "OpcaoCompra", "Classe", "Unidade"]
CAMADAS_TOT = ["TA", "BE", "CVA", "FIN"]
# monthly DMR (Diferenca Mensal de Receita) of the Tarifa Social, billed: informe de gestao, sheet 1.2
ARQ_DMR = BASE / "src" / "dmr_tarifa_social.csv"
NOTA_FORA = [
    "Retiradas da tarifa homologada (não são reembolsadas pela CDE):",
    "• Isenção de encargos (APE): com base no PRORET 5.2 (§81 e §91), o consumo isento de encargos sai do "
    "mercado de rateio da cota da CDE e não é reembolsado. Esses encargos simplesmente não são cobrados desse "
    "cliente, e o custo fica com os demais consumidores pelo rateio da cota.",
    "• SCEE - Baixa Renda: é a TE do SCEE, homologada no reajuste, aplicada sobre a energia compensada.",
]


def nota(ws, linha, coluna=1):
    """Footnote with the two parts taken out of the tariff baseline, in small grey italics."""
    for k, texto in enumerate(NOTA_FORA):
        c = ws.cell(linha + k, coluna, texto)
        c.font = openpyxl.styles.Font(size=9, italic=True, color="808080")


def le_dmr(ref):
    """Billed DMR of the month in R$, or None when the month is not in the file."""
    if not ARQ_DMR.exists():
        return None
    d = pd.read_csv(ARQ_DMR, sep=";", decimal=",", encoding="utf-8-sig", dtype={"REFERENCIA": str})
    v = d.loc[d["REFERENCIA"] == ref, "DMR_FATURADA_BRL"]
    return float(v.iloc[0]) if len(v) else None


def num(s):
    return pd.to_numeric(s, errors="coerce").fillna(0.0)


def prepara(comp, be, fator_dmr=1.0):
    """Add in-memory blocks, in the order of the bloco 2 steps: [T] = after the Tarifa Social step,
    [D] = bloco 1 minus only the DMR part of that step (fator_dmr = DMR / step), [A] = after APE,
    before the 1st pass, [P] = before the 2nd pass, [H] = tariff baseline = bloco 1 minus the parts
    not reimbursed by the CDE (SCEE - Baixa Renda = [D] - [T]; APE = [T] - [A])."""
    itens = [c[4:] for c in comp.columns if c.startswith("[1] ") and not c.startswith("[1] Total")]
    tot1_ta = num(be["[1] Total TA"])
    aj_ts = num(comp[COL_AJ_TS]) if COL_AJ_TS in comp else pd.Series(0.0, index=comp.index)
    # the Tarifa Social step is proportional to the row's bloco 1 (same ratio in every item and layer)
    r_ts = pd.Series(np.where(tot1_ta.abs() > 0.005, -aj_ts / tot1_ta.where(tot1_ta != 0, 1.0), 0.0),
                     index=comp.index)
    tot2 = num(comp["[2] Total TA"])
    aj2 = num(comp[COL_AJ_2]) if COL_AJ_2 in comp else pd.Series(0.0, index=comp.index)
    parcial = tot2 - aj2
    f2 = pd.Series(np.where((aj2 != 0) & (parcial.abs() > 0.005),
                            tot2 / parcial.where(parcial != 0, 1.0), 1.0), index=comp.index)
    for df in (comp, be):
        for i in itens:
            b1, b2 = num(df[f"[1] {i}"]), num(df[f"[2] {i}"])
            df[f"[T] {i}"] = b1 * (1.0 - r_ts)
            df[f"[D] {i}"] = b1 * (1.0 - fator_dmr * r_ts)
            df[f"[P] {i}"] = b2 / f2
            df[f"[A] {i}"] = df[f"[T] {i}"] if i in ITENS_FIO else df[f"[P] {i}"]
    # layer totals (written by the notebook in Componentes_BE for blocos 1 and 2); same in both sheets
    tot = {}
    for c in CAMADAS_TOT:
        tot[f"[1] Total {c}"] = num(be[f"[1] Total {c}"])
        tot[f"[T] Total {c}"] = tot[f"[1] Total {c}"] * (1.0 - r_ts)
        tot[f"[D] Total {c}"] = tot[f"[1] Total {c}"] * (1.0 - fator_dmr * r_ts)
        tot[f"[P] Total {c}"] = num(be[f"[2] Total {c}"]) / f2
    tot["[A] Total TA"] = comp[[f"[A] {i}" for i in itens]].sum(axis=1)
    tot["[A] Total BE"] = be[[f"[A] {i}" for i in itens]].sum(axis=1)
    # CVA + FIN of [A] is exact; its CVA/FIN split follows the row's own [T] -> [P] split
    cf_a = tot["[A] Total TA"] - tot["[A] Total BE"]
    cvat, fint = tot["[T] Total CVA"], tot["[T] Total FIN"]
    cvap, finp = tot["[P] Total CVA"], tot["[P] Total FIN"]
    d_tot = (cvat + fint) - (cvap + finp)
    share = pd.Series(np.where(d_tot.abs() > 1e-9, (cvat - cvap) / d_tot.where(d_tot != 0, 1.0), 0.0),
                      index=be.index)
    tot["[A] Total CVA"] = cvap + (cf_a - (cvap + finp)) * share
    tot["[A] Total FIN"] = cf_a - tot["[A] Total CVA"]
    for df in (comp, be):
        for k, serie in tot.items():
            df[k] = serie
    # SD rows do not go through bloco 2 (no APE, 1st or 2nd pass): the intermediate blocks are the
    # delivery itself, so any SD gap between the blocks is not attributed to a discount
    sd = (comp["Source"] == "SD").values
    for df in (comp, be):
        for i in itens:
            df.loc[sd, [f"[P] {i}", f"[A] {i}"]] = df.loc[sd, [f"[2] {i}"] * 2].values
        for c in CAMADAS_TOT:
            df.loc[sd, [f"[P] Total {c}", f"[A] Total {c}"]] = df.loc[sd, [f"[2] Total {c}"] * 2].values
    # tariff baseline [H] = [1] - ([D] - [T]) - ([T] - [A]) = [1] - [D] + [A]; SD rows keep bloco 1
    nomes = [f"{{}} {i}" for i in itens] + [f"{{}} Total {c}" for c in CAMADAS_TOT]

    def com_h(df):
        h = pd.DataFrame({k.format("[H]"): np.where(sd, df[k.format("[1]")],
                                                    df[k.format("[1]")] - df[k.format("[D]")] + df[k.format("[A]")])
                          for k in nomes}, index=df.index)
        return pd.concat([df, h], axis=1)                         # one insert: defragmented frame
    return com_h(comp), com_h(be)


def fator_dmr(ref):
    """(DMR / Tarifa Social step, step R$, DMR R$ or None); step = - "Ajuste Tarifa Social (R$)" (TA)."""
    comp = pd.read_excel(DIR_PXQ / f"pxq_{ref}.xlsx", "Componentes")
    tar = comp[comp["Tarifável?"] == "SIM"]
    estimado = float(-num(tar[COL_AJ_TS]).sum())
    dmr = le_dmr(ref)
    if dmr is None:
        print(f"ATENCAO {ref}: sem DMR em {ARQ_DMR.name}; Tarifa Social fica toda na linha DMR (estimado)")
        return 1.0, estimado, None
    return (dmr / estimado if estimado else 1.0), estimado, dmr


def layout(ref, modelo, bloco, f=1.0):
    tab, valores = grf.converte(ref, *modelo, bloco=bloco, grava=False,
                                prepara=lambda c, b: prepara(c, b, f))
    return tab, valores


def gera(ref, modelo):
    grade = [(op, cl, u) for _, op, cl, u in modelo[3]]
    f_dmr, estimado, dmr = fator_dmr(ref)
    tabs = {b: layout(ref, modelo, b, f_dmr) for b in ("[H]", "[1]", "[D]", "[T]", "[A]", "[P]", "[2]")}
    valores = tabs["[2]"][1]
    nomes_valor = [n for n in valores if n != "Total"]

    def v(bloco, k, n):
        t = tabs[bloco][0]
        x = float(t.loc[k, n]) if k in t.index else 0.0
        return x if n == "CONSUMO" else x / 1000.0               # R$ -> R$ mil

    # template: header of the delivery file of the month
    wb = openpyxl.load_workbook(DIR_RF / f"receita_faturada_{ref}.xlsx")
    ws = wb["Receita Faturada"]
    col = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value}
    estilos = [copy.copy(ws.cell(4, c)._style) for c in range(1, ws.max_column + 1)]
    competencia = ws.cell(4, col["Competência"]).value
    ws.delete_rows(4, ws.max_row)
    letra = {n: get_column_letter(c) for n, c in col.items()}
    ultima = 3 + len(grade)

    def escreve(sheet, origem):
        for i, (op, cl, u) in enumerate(grade, start=4):
            ident = dict(zip(ID_COLS, ["Faturamento", competencia, "FAT", op, cl, u]))
            for n, x in ident.items():
                sheet.cell(i, col[n], x)
            for n in nomes_valor:
                sheet.cell(i, col[n], origem((op, cl, u), n))
            sheet.cell(i, col["Total"], f"=SUM({letra['Parcela B']}{i}:{letra['Base CVA']}{i})")
            for c, st in enumerate(estilos, start=1):
                sheet.cell(i, c)._style = copy.copy(st)
            sheet.cell(i, col["Competência"]).number_format = "dd/mm/yyyy"
        t = ultima + 1
        sheet.cell(t, col["Classe"], "TOTAL").font = openpyxl.styles.Font(bold=True)
        for n in nomes_valor + ["Total"]:
            L = letra[n]
            sheet.cell(t, col[n], f"=SUM({L}4:{L}{ultima})")
            sheet.cell(t, col[n]).font = openpyxl.styles.Font(bold=True)
            sheet.cell(t, col[n]).number_format = "#,##0.0"

    ws.title = ABA_HOMOL
    escreve(ws, lambda k, n: v("[H]", k, n))
    ws.cell(2, 3, "OBSERVAÇÃO: a Tarifa Social (Baixa Renda, BPC, Indígena, Quilombola) está valorizada pela "
                  "tarifa B1 subclasse Baixa Renda homologada, sem a redução de 100% até 80 kWh (PRORET 5.2, "
                  "§129 e §130); os demais clientes, pela tarifa homologada. Sem a isenção de encargos (APE) e "
                  "sem o SCEE - Baixa Renda, que não são reembolsados pela CDE (notas abaixo). R$ mil")
    abas = {ABA_HOMOL: ws}
    for nome in [ABA_REAL, ABA_DIF] + [a for a, _, _, _ in ABAS_ORIGEM]:
        abas[nome] = wb.copy_worksheet(ws)
        abas[nome].title = nome
    escreve(abas[ABA_REAL], lambda k, n: v("[2]", k, n))
    for nome, _, de, para in ABAS_ORIGEM:
        escreve(abas[nome], lambda k, n, de=de, para=para:
                0.0 if n == "CONSUMO" else v(de, k, n) - v(para, k, n))
    for i in range(4, ultima + 2):
        for n in nomes_valor + ["Total"]:
            L = letra[n]
            abas[ABA_DIF][f"{L}{i}"] = ("=0" if n == "CONSUMO" else
                                        f"='{ABA_HOMOL}'!{L}{i}-'{ABA_REAL}'!{L}{i}")
    abas[ABA_DIF].cell(2, 3, "DIFERENÇA = tarifa homologada (bloco 1) - tarifa realizada (bloco 2); "
                             "R$ mil; positivo = desconto")
    for nome, rotulo, _, _ in ABAS_ORIGEM:
        abas[nome].cell(2, 3, f"{rotulo}: parte da diferença tarifa homologada - tarifa realizada (R$ mil)")
    nota(abas[ABA_HOMOL], ultima + 3, 2)
    nota(abas[ABA_DIF], ultima + 3, 2)
    for nome in [s for s in wb.sheetnames if s not in abas]:
        del wb[nome]

    # TA sheet: tariff of application summed (columns Parcela B .. Base CVA), one row per grid line
    tarifarias = nomes_valor[nomes_valor.index("Parcela B"):nomes_valor.index("Base CVA") + 1]

    def ta(bloco, k):
        return sum(v(bloco, k, n) for n in tarifarias)
    ws_t = wb.create_sheet(ABA_TA, index=wb.sheetnames.index(ABA_DIF) + 1)
    cab_t = (["OpcaoCompra", "Classe", "Unidade", "TA tarifa homologada", "TA tarifa realizada", "Diferença"]
             + [rot for _, rot, _, _ in ABAS_ORIGEM] + ["Soma das origens - diferença"])
    n_or = len(ABAS_ORIGEM)
    l_or1, l_orn = get_column_letter(7), get_column_letter(6 + n_or)
    ws_t.append(["Tarifa de aplicação somada (colunas I a Z do layout), R$ mil; positivo = desconto"])
    ws_t.append(cab_t)
    for i, k in enumerate(grade, start=3):
        b0, b2 = ta("[H]", k), ta("[2]", k)
        orig = [ta(de, k) - ta(para, k) for _, _, de, para in ABAS_ORIGEM]
        ws_t.append(list(k) + [b0, b2, f"=D{i}-E{i}"] + orig + [f"=SUM({l_or1}{i}:{l_orn}{i})-F{i}"])
    t = len(grade) + 3
    ws_t.cell(t, 2, "TOTAL")
    for j in range(4, 8 + n_or):
        L = get_column_letter(j)
        ws_t.cell(t, j, f"=SUM({L}3:{L}{t - 1})")
    for row in ws_t.iter_rows(min_row=2):
        for cel in row:
            if cel.column >= 4:
                cel.number_format = "#,##0.0"
            if cel.row in (2, t):
                cel.font = openpyxl.styles.Font(bold=True)
    nota(ws_t, t + 2)
    for L, w in zip("ABCDEFGHI", (14, 26, 9, 14, 14, 16, 18, 18, 22)):
        ws_t.column_dimensions[L].width = w

    # reconciliation (R$): origins vs check, and vs the totals the notebook writes
    comp, be = prepara(pd.read_excel(DIR_PXQ / f"pxq_{ref}.xlsx", "Componentes"),
                       pd.read_excel(DIR_PXQ / f"pxq_{ref}.xlsx", "Componentes_BE"), f_dmr)
    tar = comp[comp["Tarifável?"] == "SIM"]
    tot = {b: float(tabs[b][0].drop(columns="CONSUMO").values.sum()) for b in tabs}
    origem = {rot: tot[de] - tot[para] for _, rot, de, para in ABAS_ORIGEM}
    tar_be = be[be["Tarifável?"] == "SIM"]
    ts_dmr = dmr if dmr is not None else estimado
    nb_aj = {"Tarifa Social – DMR": ts_dmr, "Fonte incentivada": -num(tar[COL_AJ_1]).sum(),
             "2a passada": -num(tar[COL_AJ_2]).sum()}
    fora = {rot: tot[de] - tot[para] for rot, de, para in FORA}
    nb_fora = {"SCEE - Baixa Renda": estimado - ts_dmr, "Isenção de encargos (APE)": -num(tar.get(COL_AJ_APE, 0.0)).sum()}
    ws_c = wb.create_sheet(ABA_CONC)
    ws_c.append(["Conciliação (R$)", "Layout", "Notebook (ajustes do pxq)", "Diferença"])
    for nome in origem:
        ws_c.append([nome, origem[nome], nb_aj[nome], origem[nome] - nb_aj[nome]])
    ws_c.append(["Soma das origens", sum(origem.values()), sum(nb_aj.values()),
                 sum(origem.values()) - sum(nb_aj.values())])
    ws_c.append(["Tarifa homologada - tarifa realizada (diferença)", tot["[H]"] - tot["[2]"], None,
                 sum(origem.values()) - (tot["[H]"] - tot["[2]"])])
    ws_c.append([])
    ws_c.append(["Memória: retiradas da tarifa homologada (R$)", "Layout", "Notebook (ajustes do pxq)", "Diferença"])
    for rot in fora:
        ws_c.append([rot, fora[rot], nb_fora[rot], fora[rot] - nb_fora[rot]])
    ws_c.append(["Bloco 1 - tarifa realizada (diferença com as duas parcelas)", tot["[1]"] - tot["[2]"], None,
                 sum(origem.values()) + sum(fora.values()) - (tot["[1]"] - tot["[2]"])])
    ws_c.append([])
    ws_c.append(["Tarifa Social estimada (etapa do bloco 2, R$)", estimado])
    ws_c.append(["DMR faturada do mês (informe de gestão, R$)", dmr if dmr is not None else "sem DMR no arquivo"])
    ws_c.append(["Fator DMR / estimado", f_dmr])
    ws_c.append(["Coluna 'Notebook' das linhas da Tarifa Social: DMR do arquivo e estimado - DMR."])
    ws_c.append([])
    nota(ws_c, ws_c.max_row + 1)
    ws_c.append([])
    ws_c.append([])
    ws_c.append([])
    ws_c.append(["Positivo = desconto (bloco 1 maior que o bloco 2). A 2a passada não é desconto."])
    ws_c.append(["CVA e Base Financeira: divisão entre APE e fonte incentivada proporcional na linha "
                 "do pxq que tem as duas etapas."])
    for row in ws_c.iter_rows():
        for cel in row:
            if isinstance(cel.value, float):
                cel.number_format = "#,##0.00"
    ws_c.column_dimensions["A"].width = 40
    for L in "BCD":
        ws_c.column_dimensions[L].width = 22

    DIR_SAIDA.mkdir(parents=True, exist_ok=True)
    saida = DIR_SAIDA / f"receita_faturada_Bl1_vs_Bl2_{ref}.xlsx"
    wb.save(saida)
    print(f"{ref}: homologada - realizada R$ {(tot['[H]'] - tot['[2]'])/1e6:,.2f} mi | Tarifa Social DMR "
          f"{origem['Tarifa Social – DMR']/1e6:,.2f} | fonte incentivada {origem['Fonte incentivada']/1e6:,.2f} | "
          f"2a passada {origem['2a passada']/1e6:,.2f} | fora: SCEE - Baixa Renda "
          f"{fora['SCEE - Baixa Renda']/1e6:,.2f}, APE {fora['Isenção de encargos (APE)']/1e6:,.2f} | {saida.name}")


def consolida(refs):
    """One row per month (sum of every layout row) in each sheet."""
    out = openpyxl.Workbook()
    out.remove(out.active)
    dados, nomes = {}, None
    for ref in refs:
        wb = openpyxl.load_workbook(DIR_SAIDA / f"receita_faturada_Bl1_vs_Bl2_{ref}.xlsx", read_only=True)
        for aba in [ABA_HOMOL, ABA_REAL] + [a for a, _, _, _ in ABAS_ORIGEM]:
            rows = list(wb[aba].iter_rows(min_row=3, values_only=True))
            cab = list(rows[0])
            idx = [i for i, n in enumerate(cab) if n and n not in ID_COLS + ["Total"]]
            nomes = [cab[i] for i in idx]
            corpo = [r for r in rows[1:] if r[cab.index("OpcaoCompra")]]
            dados.setdefault(aba, []).append([ref] + [sum(float(r[i] or 0.0) for r in corpo
                                                          if isinstance(r[i], (int, float))) for i in idx])
    cab = ["Mês"] + nomes + ["Total I a Z", "Total geral"]
    n = len(nomes)
    l_ini, l_z, l_ult = (get_column_letter(cab.index(x) + 1) for x in ("Parcela B", "Base CVA", nomes[-1]))
    for aba in [ABA_HOMOL, ABA_REAL, ABA_DIF] + [a for a, _, _, _ in ABAS_ORIGEM]:
        ws = out.create_sheet(aba)
        ws.append(cab)
        for k, ref in enumerate(refs, start=2):
            ws.cell(k, 1, ref)
            for j in range(n):
                c = get_column_letter(j + 2)
                ws.cell(k, j + 2, f"='{ABA_HOMOL}'!{c}{k}-'{ABA_REAL}'!{c}{k}" if aba == ABA_DIF
                        else dados[aba][k - 2][j + 1])
                ws.cell(k, j + 2).number_format = "#,##0.0"
            ws.cell(k, n + 2, f"=SUM({l_ini}{k}:{l_z}{k})").number_format = "#,##0.0"
            ws.cell(k, n + 3, f"=SUM({l_ini}{k}:{l_ult}{k})").number_format = "#,##0.0"
        for c in ws[1]:
            c.font = openpyxl.styles.Font(bold=True)
        ws.cell(len(refs) + 3, 1, "R$ mil; positivo = desconto; consumo em MW + MWh")
        if aba in (ABA_HOMOL, ABA_DIF):
            nota(ws, len(refs) + 4)
    saida = DIR_SAIDA / f"consolidado_receita_faturada_Bl1_vs_Bl2_{refs[0]}_a_{refs[-1]}.xlsx"
    out.save(saida)
    print(f"consolidado: {saida.name}")


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    if args == ["--todos"]:
        args = sorted(p.stem[4:] for p in DIR_PXQ.glob("pxq_????-??.xlsx"))
    elif len(args) == 2:
        args = [p.strftime("%Y-%m") for p in pd.period_range(args[0], args[1], freq="M")]
    modelo = grf.le_modelo()
    for ref in args:
        gera(ref, modelo)
    if len(args) > 1:
        consolida(args)
