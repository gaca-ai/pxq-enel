"""
Converte a saida do PxQ (1-Entrega/pxq/pxq_<ref>.xlsx) no modelo Receita Faturada
solicitado pela area (tbl/modelo/De-Para_Modelo_PyC.xlsx, abas "Receita Faturada" e "Instrucoes").

Uso:
    python gera_receita_faturada.py 2026-06            # um mes
    python gera_receita_faturada.py 2026-01 2026-08    # varios meses
    python gera_receita_faturada.py --todos            # todos os pxq_<ref>.xlsx

Regras (README, secao 5.1):
- Colunas I a X: bloco 2 na camada BE, componente a componente, pelo de-para da
  aba "Instrucoes" do proprio modelo (colunas L e N). Y = Base Financeira (total FIN),
  Z = Base CVA (total CVA). H = quantidade em MWh (Unidade E) ou MW (Unidade D).
- TipoFaturamento pela coluna "Tipo de faturamento" do pxq: CNR (TOI, Faturamento Manual de
  Energia), REFAT (flag de refaturamento X) e FAT (demais e todo o SD). Cada tipo tem a propria
  grade de 77 linhas no modelo, na ordem FAT, REFAT, CNR. Hibrido entra como LIVRE. Baixa Renda reune Baixa Renda,
  BPC, Indigena, Quilombola e Desconto Social.
- AA (Bandeira): Grp.estatists.quant. nos 8 codigos de bandeira, unidade E.
- AB (PIS/COFINS) e AC (ICMS): Operacao principal R/#. Parcela TE vai para a linha E;
  parcela TUSD e os sem parcela (inclusive subvencoes) sao repartidos entre D e E na
  proporcao da receita TUSD (ou total) do bloco 2 da mesma opcao x classe e tipo de faturamento
  (sem bloco 2 no tipo, vale a proporcao de todos os tipos juntos).
- ZQC237 (Encargo da conta COVID) entra em CDE CONTAS, unidade E.
- UD/ER (coluna nova, antes de PIS/COFINS): ultrapassagem de demanda e energia reativa
  (ultrapassagem, UFER, UFDR, demanda e energia reativa excedente).
- Outros (coluna nova, antes de PIS/COFINS): os demais nao tarifaveis (acertos da CCEE,
  energia compensada, atualizacao de montante, ERCAP, desconto ex-funcionario, encargo de
  conexao, substituicao tributaria - debito e credito se anulam por classe) e o residuo
  sem item tarifario do bloco 2.
- Unidade de UD/ER e Outros pelo proprio lancamento: demanda, potencia, kW, UFDR e conexao
  vao para D; o restante para E.
- Com isso o layout reproduz a receita total do pxq (BW + SD), a base do informe de gestao.
- SCEE GD1/GD2 ficam vazios: o BW nao identifica a instalacao no SCEE; o movimento
  desses clientes fica em CATIVO (README 6.1).
- Classe principal "Nao atribuido": categoria residencial vai para Residencial, ou para
  Residencial Baixa Renda quando todas as subclasses listadas sao de baixa renda ou
  desconto social.
- Linha E sem correspondente no modelo: o valor vai para Outros da linha D da mesma opcao; o MWh
  nao e gravado e aparece em nota na aba Conferencia. O modelo usado (De-Para_Modelo_PyC_sugerido.xlsx)
  tem a linha GERACAO / E, e o caso deixou de ocorrer com a energia das geradoras.
- PIS/COFINS e ultrapassagem do SD (geradoras, suprimentos e uso da distribuicao) vem do pxq como
  nao tarifaveis e vao para PIS/COFINS e UD/ER da linha da propria opcao.
- DEMANDA ULTRAPASSAGEM (coluna nova, logo apos CONSUMO, em MW, so na linha D): demanda de
  ultrapassagem dos lancamentos nao tarifaveis de ultrapassagem (Grp.estatists.mont. com "ULTRAPA"),
  coluna "Demanda (kW)" da aba Nao_tarifaveis_linha do pxq. Cativo (ZQD021 a ZQD023) e SD trazem a quantidade; a ultrapassagem
  do livre (ZQD019, ZQD020, ZQD037) vem do BW so com R$, sem kW. E quantidade: nao entra no Total
  nem na conciliacao de R$, e a demanda de ultrapassagem nao esta no CONSUMO.
- Valores em R$ mil; so as linhas com valor sao gravadas.
"""
import os
import copy
import sys
import datetime as dt
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl

BASE = Path(os.environ.get("PXQ_BASE", r"C:\python_projects\claude\p_x_q\SP\BW"))
# Layout model: the area's model plus a GERACAO / E row (energy of the generators, sessions P27),
# kept in a separate file so the area's model stays untouched. To go back, point to
# "De-Para_Modelo_PyC.xlsx".
ARQ_MODELO = BASE / "tbl" / "modelo" / "De-Para_Modelo_PyC_sugerido.xlsx"   # versioned
DIR_PXQ = BASE / "1-Entrega" / "pxq"
DIR_SAIDA = BASE / "1-Entrega" / "receita_faturada"

CODIGOS_BANDEIRA = {"ZQDFAM", "ZQDPAM", "ZQDFVM", "ZQDPVM", "ZQODAM", "ZQODVM", "ZQDIAM", "ZQDIVM"}
CODIGO_COVID = "ZQC237"
SUBCLASSES_BAIXA_RENDA = ("Baixa Renda", "Desconto Social")

OPCAO = {"Cliente Cativo": "CATIVO", "Cliente Livre": "LIVRE", "Cliente Híbrido": "LIVRE",
         "GERAÇÃO": "GERAÇÃO", "SUPRIMENTO": "SUPRIMENTO", "USO DISTRIBUIÇÃO": "USO D"}
CLASSE = {"Residencial": "Residencial", "Industrial": "Industrial", "Comercial, Serviços": "Comercial",
          "Rural": "Rural", "Poder Público": "Poderes Públicos", "Iluminação Pública": "Iluminação Pública",
          "Serviço Público": "Serviços Públicos", "Consumo Próprio": "Consumo Próprio"}
OPCOES_SD = {"GERAÇÃO", "SUPRIMENTO", "USO D"}                     # SD options: class = option
COL_ID = ["Responsável", "Competência", "TipoFaturamento", "OpcaoCompra", "Classe", "Unidade"]
TIPOS = ("FAT", "REFAT", "CNR")                                    # billing types, in the model order
COL_TIPO = "Tipo de faturamento"
COL_UDER, COL_OUTROS = "UD/ER", "Outros"                          # new columns, before PIS/COFINS
COL_UD_MW = "DEMANDA ULTRAPASSAGEM"                                # new quantity column (MW), after CONSUMO
PADRAO_UD_MW = "ULTRAPA"                                           # demand overrun lines (non-tariff)
PADRAO_UDER = "ULTRAPA|UFER|UFDR|REATIV"
PADRAO_UNID_D = "DEMANDA|POTÊNCIA|POTENCIA|\\(KW\\)|UFDR|CONEX"


def norm(texto):
    """Case- and accent-insensitive key for matching model labels."""
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(t.upper().split())


def le_modelo():
    """Template workbook, header row (row 3) and component -> model column from 'Instruções'."""
    wb = openpyxl.load_workbook(ARQ_MODELO)
    ws = wb["Receita Faturada"]
    cab = {norm(ws.cell(3, c).value): ws.cell(3, c).value
           for c in range(2, ws.max_column + 1) if ws.cell(3, c).value}
    ins = wb["Instruções"]
    de_para, grupo = {}, None
    for r in range(12, ins.max_row + 1):
        rotulo, comp = ins.cell(r, 12).value, ins.cell(r, 14).value
        if rotulo:
            grupo = rotulo
        if comp and grupo and norm(grupo) in cab and "TOTAL" not in norm(comp):
            de_para[str(comp).strip()] = cab[norm(grupo)]
    # template FAT grid, in the model's own order
    grade = [tuple(ws.cell(r, c).value for c in range(4, 8))
             for r in range(4, ws.max_row + 1) if ws.cell(r, 4).value == "FAT"]
    return wb, cab, de_para, grade


def classe_modelo(classe_principal, subclasse, categoria=None):
    # "Nao atribuido" main class: residential tariff category -> Residencial, or Baixa Renda when
    # every listed subclass is a low-income / social-discount one
    if norm(classe_principal) == "NAO ATRIBUIDO":
        if "RESIDENCIAL" not in norm(categoria):
            return None
        partes = [p.strip() for p in str(subclasse).split("|") if p.strip() and p.strip() != "nan"]
        so_br = bool(partes) and all(any(b in p for b in SUBCLASSES_BAIXA_RENDA) for p in partes)
        return "Residencial Baixa Renda" if so_br else "Residencial"
    if classe_principal == "Residencial" and any(s in str(subclasse) for s in SUBCLASSES_BAIXA_RENDA):
        return "Residencial Baixa Renda"
    return CLASSE.get(classe_principal)


def converte(ref, wb, cab, de_para, grade, bloco="[2]", grava=True, prepara=None):
    """Layout of month `ref` from the pxq, using the component columns of `bloco` ("[2]" = the
    delivery). `prepara(comp, be)` may add in-memory blocks; with grava=False nothing is written and
    the layout table (R$, indexed by option x class x unit, billing types summed) is returned."""
    arq = DIR_PXQ / f"pxq_{ref}.xlsx"
    comp = pd.read_excel(arq, "Componentes")
    be = pd.read_excel(arq, "Componentes_BE")
    nt = pd.read_excel(arq, "Nao_tarifaveis_linha")
    if prepara is not None:
        comp, be = prepara(comp, be)
    pre = f"{bloco} "
    receita_total = comp["RECEITA DE FORNECIMENTO"].sum()

    # ------------------------------------------------ tariff groups (BW + SD), BE layer
    tar = be[be["Tarifável?"] == "SIM"].copy()
    tar_ta = comp.loc[tar.index]                                   # same row order in both sheets
    tar["OPCAO"] = tar["Flag Cli.Liv./ Cat."].map(OPCAO)
    tar["TIPO"] = tar[COL_TIPO].fillna("FAT") if COL_TIPO in tar else "FAT"
    sd = tar["Source"] == "SD"
    tar["CLASSE"] = [o if s else classe_modelo(c, sc, ct) for o, s, c, sc, ct in
                     zip(tar["OPCAO"], sd, tar["Classe principal"], tar["Classe/Sub-classe"],
                         tar["Categoria tarifa"])]
    tar["UNIDADE"] = np.where(tar["Chave"].astype(str).str.endswith("kW"), "D", "E")
    tar["CONSUMO"] = np.where(tar["UNIDADE"] == "D", tar["TOTAL DEMANDA"], tar["TOTAL CONSUMO"]) / 1000
    cols_comp = [c for c in be.columns if c.startswith(pre) and not c.startswith(pre + "Total")]
    # SD keys valued by the SD own tariff sheet carry only the TA layer (no BE/CVA/FIN split):
    # their TA components go to the economic columns, and the row is listed in Conferencia
    so_ta = (tar["[2] Total TA"].fillna(0.0) - tar[["[2] Total BE", "[2] Total CVA", "[2] Total FIN"]]
             .fillna(0.0).sum(axis=1)).abs() > 0.01
    tar.loc[so_ta, cols_comp] = tar_ta.loc[so_ta, cols_comp].values
    tar.loc[so_ta, pre + "Total BE"] = tar.loc[so_ta, pre + "Total TA"]
    notas = [f"Tarifa própria do SD sem camadas (só TA), alocada na base econômica: "
             f"{r['Chave']} - R$ {r['[2] Total TA']:,.2f}" for _, r in tar[so_ta].iterrows()]
    for c in cols_comp:
        destino = de_para.get(c[len(pre):])
        if destino:
            tar[destino] = tar.get(destino, 0.0) + tar[c].fillna(0.0)
    sem_coluna = tar[[c for c in cols_comp if c[len(pre):] not in de_para]].fillna(0.0).sum(axis=1)
    tar[cab["BASE FINANCEIRA"]] = tar[pre + "Total FIN"].fillna(0.0)
    tar[cab["BASE CVA"]] = tar[pre + "Total CVA"].fillna(0.0)
    tar["RESIDUO"] = tar["RECEITA DE FORNECIMENTO"] - tar["[2] Total TA"].fillna(0.0)
    # TUSD and total TA by unit, used to split taxes between D and E
    tusd = [c for c in comp.columns if c.startswith("[2] TUSD")]
    tar["PESO_TUSD"] = tar_ta[tusd].fillna(0.0).sum(axis=1)
    tar["PESO_TOTAL"] = tar_ta["[2] Total TA"].fillna(0.0)

    valores = [c for c in cab.values() if c not in COL_ID]
    pos = valores.index(cab["PIS/COFINS"])
    valores[pos:pos] = [COL_UDER, COL_OUTROS]
    valores.insert(valores.index(cab["CONSUMO"]) + 1, COL_UD_MW)
    qtd = {cab["CONSUMO"], COL_UD_MW}                              # quantity columns (not R$)
    linhas = []
    fora = []                                                      # (category, detail, R$)
    ok_tar = tar["OPCAO"].notna() & tar["CLASSE"].notna()
    for (o, cl), r in tar[~ok_tar].groupby(["Flag Cli.Liv./ Cat.", "Classe principal"], dropna=False):
        fora.append(("Classe/opcao sem correspondencia no modelo", f"{o} / {cl}",
                     r["RECEITA DE FORNECIMENTO"].sum()))
    for tp, r in tar[ok_tar & ~tar["TIPO"].isin(TIPOS)].groupby("TIPO"):
        fora.append(("Tipo de faturamento sem grade no modelo", tp, r["RECEITA DE FORNECIMENTO"].sum()))
    t = tar[ok_tar & tar["TIPO"].isin(TIPOS)]
    blocos = [c for c in valores if c in t.columns]
    chave = ["TIPO", "OPCAO", "CLASSE", "UNIDADE"]
    agg = t.groupby(chave)[blocos + ["PESO_TUSD", "PESO_TOTAL"]].sum()
    residuo = t.groupby(chave)["RESIDUO"].sum()
    peso_todos = agg.groupby(level=[1, 2, 3])[["PESO_TUSD", "PESO_TOTAL"]].sum()   # all types together

    # ------------------------------------------------ non-tariff lines
    nt = nt.copy()
    nt["OPCAO"] = nt["Flag Cli.Liv./ Cat."].map(OPCAO)
    nt["TIPO"] = nt[COL_TIPO].fillna("FAT") if COL_TIPO in nt else "FAT"
    nt["CLASSE"] = [o if o in OPCOES_SD else classe_modelo(c, s, ct) for o, c, s, ct in
                    zip(nt["OPCAO"], nt["Classe principal"], nt["Classe/Sub-classe"], nt["Categoria tarifa"])]
    mont = nt["Grp.estatísts.mont."].astype(str)
    quant = nt["Grp.estatísts.quant."].astype(str)
    tributo = nt["Operação principal"].astype(str).str.contains("R/#", regex=False)
    nt["COLUNA"] = None
    nt.loc[quant.isin(CODIGOS_BANDEIRA), "COLUNA"] = cab["BANDEIRA"]
    nt.loc[quant == CODIGO_COVID, "COLUNA"] = cab["CDE CONTAS"]
    nt.loc[tributo & mont.str.contains("ICMS"), "COLUNA"] = cab["ICMS"]
    nt.loc[tributo & mont.str.contains("PIS|COFINS"), "COLUNA"] = cab["PIS/COFINS"]
    nt["PARCELA"] = np.where(mont.str.endswith("- TE"), "TE",
                             np.where(mont.str.endswith("- TUSD"), "TUSD", "GERAL"))
    nt.loc[nt["COLUNA"].isin([cab["BANDEIRA"], cab["CDE CONTAS"]]), "PARCELA"] = "TE"
    mont_up = mont.str.upper()
    nt.loc[nt["COLUNA"].isna() & mont_up.str.contains(PADRAO_UDER), "COLUNA"] = COL_UDER
    nt.loc[nt["COLUNA"].isna(), "COLUNA"] = COL_OUTROS
    novas = nt["COLUNA"].isin([COL_UDER, COL_OUTROS])
    nt.loc[novas, "PARCELA"] = np.where(mont_up[novas].str.contains(PADRAO_UNID_D), "UNID_D", "UNID_E")

    usa = nt["COLUNA"].notna() & nt["OPCAO"].notna() & nt["CLASSE"].notna()
    for g, r in nt[nt["COLUNA"].notna() & ~usa].groupby(["Flag Cli.Liv./ Cat.", "Classe principal"], dropna=False):
        fora.append(("Nao tarifavel com classe sem correspondencia", " / ".join(map(str, g)),
                     r["RECEITA DE FORNECIMENTO"].sum()))
    for tp, r in nt[usa & ~nt["TIPO"].isin(TIPOS)].groupby("TIPO"):
        fora.append(("Tipo de faturamento sem grade no modelo", tp, r["RECEITA DE FORNECIMENTO"].sum()))
    usa &= nt["TIPO"].isin(TIPOS)

    extra = {}
    for (tp, o, cl, col, parc), v in (nt[usa].groupby(["TIPO", "OPCAO", "CLASSE", "COLUNA", "PARCELA"])
                                      ["RECEITA DE FORNECIMENTO"].sum().items()):
        if parc == "TE" or parc == "UNID_E":
            partes = {"E": 1.0}
        elif parc == "UNID_D":
            partes = {"D": 1.0}
        else:
            peso = "PESO_TUSD" if parc == "TUSD" else "PESO_TOTAL"
            w = {u: agg.loc[(tp, o, cl, u), peso] if (tp, o, cl, u) in agg.index else 0.0 for u in ("D", "E")}
            if not w["D"] + w["E"]:                                 # no tariff block in this type
                w = {u: peso_todos.loc[(o, cl, u), peso] if (o, cl, u) in peso_todos.index else 0.0
                     for u in ("D", "E")}
            tot = w["D"] + w["E"]
            partes = {u: w[u] / tot for u in w} if tot else {"E": 1.0}
        for u, p in partes.items():
            extra[(tp, o, cl, u, col)] = extra.get((tp, o, cl, u, col), 0.0) + v * p
    for (tp, o, cl, u), v in residuo.items():
        if abs(v) > 0.005:
            extra[(tp, o, cl, u, COL_OUTROS)] = extra.get((tp, o, cl, u, COL_OUTROS), 0.0) + v

    # demand overrun quantity (MW), D row of the option x class: overrun lines of Nao_tarifaveis_linha
    # (one row per BW line, so the class is exact); pxq generated before the "Demanda (kW)" column
    # there: the column stays empty, with a warning
    if "Demanda (kW)" in nt:
        ud = nt[usa & mont_up.str.contains(PADRAO_UD_MW)]
        for (tp, o, cl), v in (ud.groupby(["TIPO", "OPCAO", "CLASSE"])["Demanda (kW)"].sum() / 1000).items():
            if abs(v) > 1e-9:
                extra[(tp, o, cl, "D", COL_UD_MW)] = extra.get((tp, o, cl, "D", COL_UD_MW), 0.0) + v
    else:
        print(f"ATENCAO {ref}: pxq sem a coluna 'Demanda (kW)' em Nao_tarifaveis_linha; "
              f"{COL_UD_MW} fica vazia (rodar o notebook de novo)")

    # ------------------------------------------------ assemble rows on the model grid
    extra_df = (pd.Series(extra, dtype=float).rename_axis(chave + ["COLUNA"])
                .unstack("COLUNA") if extra else pd.DataFrame())
    tabela = (pd.concat([agg[blocos], extra_df]).groupby(level=[0, 1, 2, 3]).sum()
              .reindex(columns=valores).fillna(0.0))

    competencia = dt.datetime(int(ref[:4]), int(ref[5:7]), 1)
    na_grade = {(op, cl, u) for _, op, cl, u in grade}
    for k, r in list(tabela.iterrows()):
        tp, k = k[0], k[1:]
        if k in na_grade:
            continue
        valor = r.drop(list(qtd)).sum()
        alvo = (k[0], k[1], "D")
        if k[2] == "E" and alvo in na_grade:
            # no E row in the model for this option (GERACAO): value goes to Outros of the D row;
            # the MWh cannot be added to MW and is only reported
            if (tp,) + alvo not in tabela.index:
                tabela.loc[(tp,) + alvo, :] = 0.0
            tabela.loc[(tp,) + alvo, COL_OUTROS] += valor
            notas.append(f"{tp} / {' / '.join(k)} sem linha E no modelo: R$ {valor:,.2f} gravados em "
                         f"Outros da linha D ({r[cab['CONSUMO']]:,.3f} MWh não gravados)")
        else:
            fora.append(("Opcao x classe x unidade fora da grade do modelo", " / ".join((tp,) + k), valor))
    for tp in TIPOS:
        for _, op, cl, u in grade:
            if (tp, op, cl, u) in tabela.index:
                r = tabela.loc[(tp, op, cl, u)]
                if r.abs().sum() > 0.005:
                    linhas.append(["Faturamento", competencia, tp, op, cl, u] + list(r.values))
    if not grava:
        soma = tabela.groupby(level=[1, 2, 3]).sum()                # billing types summed
        return soma[[k in na_grade for k in soma.index]], valores

    # ------------------------------------------------ write output from the template
    wb = openpyxl.load_workbook(ARQ_MODELO)
    ws = wb["Receita Faturada"]
    larguras = {c: ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width
                for c in range(1, ws.max_column + 1)}
    c_pis = next(c for c in range(2, ws.max_column + 1) if ws.cell(3, c).value == cab["PIS/COFINS"])
    ws.insert_cols(c_pis, 2)
    for k, nome in enumerate([COL_UDER, COL_OUTROS]):
        for r in (2, 3):
            ws.cell(r, c_pis + k).value = ws.cell(r, c_pis + 2).value if r == 2 else nome
            ws.cell(r, c_pis + k)._style = copy.copy(ws.cell(r, c_pis + 2)._style)
        ws.cell(4, c_pis + k)._style = copy.copy(ws.cell(4, c_pis + 2)._style)
    for c, w in larguras.items():                                 # shift the column widths
        if w:
            ws.column_dimensions[openpyxl.utils.get_column_letter(c if c < c_pis else c + 2)].width = w
    ws.column_dimensions[openpyxl.utils.get_column_letter(c_pis)].width = 10
    ws.column_dimensions[openpyxl.utils.get_column_letter(c_pis + 1)].width = 10
    # demand overrun quantity column right after CONSUMO, same style as CONSUMO
    c_cons = next(c for c in range(2, ws.max_column + 1) if ws.cell(3, c).value == cab["CONSUMO"])
    larguras = {c: ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width
                for c in range(1, ws.max_column + 1)}
    ws.insert_cols(c_cons + 1, 1)
    for r in (2, 3, 4):
        ws.cell(r, c_cons + 1)._style = copy.copy(ws.cell(r, c_cons)._style)
    ws.cell(2, c_cons + 1).value = "MW"
    ws.cell(3, c_cons + 1).value = COL_UD_MW
    for c, w in larguras.items():                                 # shift the column widths
        if w and c > c_cons:
            ws.column_dimensions[openpyxl.utils.get_column_letter(c + 1)].width = w
    ws.column_dimensions[openpyxl.utils.get_column_letter(c_cons + 1)].width = 14
    estilos = [copy.copy(ws.cell(4, c)._style) for c in range(1, ws.max_column + 1)]
    ws.delete_rows(4, ws.max_row)
    col_de = {ws.cell(3, c).value: c for c in range(2, ws.max_column + 1)}
    letra_i, letra_z = (openpyxl.utils.get_column_letter(col_de[cab[k]]) for k in ("PARCELA B", "BASE CVA"))
    for i, lin in enumerate(linhas, start=4):
        for nome, v in zip(COL_ID + valores, lin):
            if nome not in qtd and nome not in COL_ID:
                v = v / 1000                                         # R$ -> R$ mil
            ws.cell(i, col_de[nome], v)
        ws.cell(i, col_de[cab["TOTAL"]], f"=SUM({letra_i}{i}:{letra_z}{i})")
        for c, st in enumerate(estilos, start=1):
            ws.cell(i, c)._style = copy.copy(st)
        ws.cell(i, col_de["Competência"]).number_format = "dd/mm/yyyy"

    # ------------------------------------------------ reconciliation sheet (R$)
    fora.append(("Componente BE sem coluna no modelo", "", sem_coluna.sum()))
    df_fora = pd.DataFrame(fora, columns=["Categoria", "Detalhe", "R$"])
    df_fora = df_fora[df_fora["R$"].abs() > 0.005]
    i_valor = [6 + k for k, n in enumerate(valores) if n not in qtd]   # R$ positions in each row
    gravado = sum(sum(l[i] for i in i_valor) for l in linhas)
    resumo = pd.DataFrame([
        ("Receita total do pxq (BW + SD)", receita_total),
        ("Gravado no layout (colunas de valor)", gravado),
        ("Fora do layout (detalhe abaixo)", df_fora["R$"].sum()),
        ("Diferenca (deve ser zero)", receita_total - gravado - df_fora["R$"].sum()),
    ], columns=["Item", "R$"])
    if abs(resumo["R$"].iloc[-1]) > 1.0:
        print(f"ATENCAO {ref}: conciliacao nao fecha em R$ {resumo['R$'].iloc[-1]:,.2f}")
    ws_c = wb.create_sheet("Conferencia")
    ws_c.append(["Conciliação (R$)", None])
    for r in resumo.itertuples(index=False):
        ws_c.append(list(r))
    ws_c.append([])
    ws_c.append(["Categoria fora do layout", "Detalhe", "R$"])
    for r in df_fora.sort_values(["Categoria", "R$"]).itertuples(index=False):
        ws_c.append(list(r))
    ws_c.append([])
    ws_c.append(["Gravado por tipo de faturamento (R$)", None])
    for tp in TIPOS:
        ws_c.append([tp, sum(sum(l[i] for i in i_valor) for l in linhas if l[2] == tp)])
    ws_c.append([])
    for nota in notas:
        ws_c.append([nota])
    ws_c.append(["Pendência: SCEE - GD1 / SCEE - GD2 não são separados (o BW não identifica a instalação no SCEE; README 6.1); "
                 "o movimento desses clientes está em CATIVO."])
    for row in ws_c.iter_rows():
        for cel in row:
            if isinstance(cel.value, float):
                cel.number_format = "#,##0.00"
    ws_c.column_dimensions["A"].width = 62
    ws_c.column_dimensions["B"].width = 48
    ws_c.column_dimensions["C"].width = 18

    # ------------------------------------------------ BW lots outside the reading calendar de-para
    # (copied from the pxq: FAT lines of these lots used the month tariff, without proration)
    ws_l = wb.create_sheet("Lotes sem calendário")
    try:
        lotes = pd.read_excel(arq, "Lotes_sem_calendario")
    except ValueError:                       # pxq generated before the reading calendar
        lotes = None
    if lotes is None:
        ws_l.append(["Aba Lotes_sem_calendario ausente no pxq (gerado antes do calendário de leitura)."])
    elif lotes.empty:
        ws_l.append(["Nenhum lote do BW fora do de-para do calendário de leitura (tbl/calendario/de_para_lote.csv)."])
    else:
        ws_l.append(list(lotes.columns))
        for r in lotes.itertuples(index=False):
            ws_l.append(list(r))
        for cel in ws_l["C"][1:]:
            cel.number_format = "#,##0.00"
        ws_l.append([])
        ws_l.append(["Lotes com faturamento normal (FAT) fora do de-para do calendário de leitura: ficam com a "
                     "tarifa do mês, sem tarifa proporcional. Incluir o lote em tbl/calendario/de_para_lote.csv."])
    ws_l.column_dimensions["A"].width = 40
    ws_l.column_dimensions["B"].width = 20
    ws_l.column_dimensions["C"].width = 18

    DIR_SAIDA.mkdir(parents=True, exist_ok=True)
    saida = DIR_SAIDA / f"receita_faturada_{ref}.xlsx"
    wb.save(saida)
    por_tipo = " | ".join(f"{tp} {sum(sum(l[i] for i in i_valor) for l in linhas if l[2] == tp)/1e6:,.2f}"
                         for tp in TIPOS)
    print(f"{ref}: {len(linhas)} linhas ({por_tipo}) | gravado R$ {gravado/1e6:,.2f} mi | fora R$ "
          f"{df_fora['R$'].sum()/1e6:,.2f} mi | diferenca R$ {resumo['R$'].iloc[-1]:,.2f} | {saida.name}")


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    if args == ["--todos"]:
        args = sorted(p.stem[4:] for p in DIR_PXQ.glob("pxq_????-??.xlsx"))
    modelo = le_modelo()
    for ref in args:
        converte(ref, *modelo)
