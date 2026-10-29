# -*- coding: utf-8 -*-
"""Extract the homologated tariff (TA, BE, CVA, FIN) from the PCAT workbooks and
build ONE long table, one row per (vigencia x camada x chave), with
VIGENCIA_INICIO / VIGENCIA_FIM as real dates.

This replaces the previous design (a variable per vintage -- ARQ_TARIFA_ATUAL,
ARQ_TARIFA_ANTERIOR -- that had to be renamed by hand every year). Adding a new
vintage from here on means adding one entry to VIGENCIAS below and re-running
this script; nothing else in the notebook changes.

Columns are matched by NAME (header row, row 3 of each PCAT sheet), never by
fixed position, so a component added or removed by ANEEL between vintages
(e.g. TE_ANGRA, introduced in the 2026 vintage) does not break the read --
it is reported below and filled with 0.0 in the vintages where it is absent.

Output: tbl/tarifas/tarifa_long.csv
"""
import os
import datetime as dt
from pathlib import Path

import pandas as pd
import openpyxl

BASE = Path(os.environ.get("PXQ_BASE", r"C:\python_projects\claude\p_x_q\SP\BW"))
OUTDIR = BASE / "tbl" / "tarifas"

# One entry per PCAT vintage. Add a line here when a new vintage is published;
# nothing else needs to change.
VIGENCIAS = [
    {"ano": 2024, "arquivo": BASE / "src" / "pcats" / "390 PCAT Enel SP 2024 V02.xlsx",
     "inicio": dt.date(2024, 7, 4), "fim": dt.date(2025, 7, 3)},
    {"ano": 2025, "arquivo": BASE / "src" / "pcats" / "390 PCAT Enel SP 2025 V02_25jun TA.xlsx",
     "inicio": dt.date(2025, 7, 4), "fim": dt.date(2026, 7, 3)},
    {"ano": 2026, "arquivo": BASE / "src" / "pcats" / "390 PCAT Enel SP 2026 V02_23junho.xlsx",
     "inicio": dt.date(2026, 7, 4), "fim": dt.date(2027, 7, 3)},
]

ABAS = {"TA": "TA - Aplicação", "BE": "TA - BE", "CVA": "TA - CVA"}
ATRIBUTOS = ["SUBGRUPO", "MODALIDADE", "CLASSE", "SUBCLASSE", "DETALHE",
             "POSTO", "UNIDADE", "ACESSANTE"]
TOTAIS = ["Total TUSD", "Total TE"]

# Desconto Social (Lei 15.235/2025, REN ANEEL 1147/2025). PCATs published before
# the benefit existed (2024, 2025) have no row for it, but it is billed since
# 2026. The tariff is fully defined by the regulation, so the missing rows are
# derived from the B1 Residencial row of the same vintage:
#   faixa 01 (up to 120 kWh): Residencial without the exempt items
#            (PRORET 7.1 par. 18.V - TUSD alineas d/f; par. 23.V - TE alinea f);
#   faixa 02 (above 120 kWh): Residencial as is (REN 1147, art. 179 par. 4 II).
# Same rule the Mercado Receita applies (checked on mar/2026: 306,15 / 269,02).
# Derived rows carry ORIGEM = "DERIVADA_RESIDENCIAL"; PCAT rows carry "PCAT".
# In vintages that already publish the benefit the native rows are kept and the
# derivation only runs as a check against them.
DS_FAIXA_01 = "Residencial Desconto Social – faixa 01"
DS_FAIXA_02 = "Residencial Desconto Social – faixa 02"
ISENTOS_DESCONTO_SOCIAL = ["TUSD_CDE", "TUSD_CDE_COVID", "TE_CDE_GD"]


# TE of the supply to distributors (A4 Distribuição keys: Ceris, CPFL Piratininga, Neoenergia
# Elektro). Vintages before 2026 publish these keys with TE = 0 in the "TA - ..." sheets; the TE
# of supply is only in the TE sheets, row "Energia suprimento permissionária / concessionária"
# (<500GWh), with the layers apart: "TE BE" (BE), "TE CVA" (CVA) and "TE BF" (financeiros, FIN).
# Those rows fill the TE items of the keys whose TE is zero; ORIGEM = "TE_SUPRIMENTO".
# (sessions P1: the SD bills the Ceris TE at exactly this tariff, 278,27 R$/MWh in 2025.)
ABAS_TE_SUPRIMENTO = {"BE": "TE BE", "CVA": "TE CVA", "FIN": "TE BF"}
LINHA_SUPRIMENTO = {"Ceris": "Energia suprimento permissionária"}   # other acessantes: concessionária
LINHA_SUPRIMENTO_PADRAO = "Energia suprimento concessionária"
# item label in the TE sheets (row 4) -> item column of the tariff table
ITEM_TE = {
    "P&D": "TE_PeD", "ESS/ERR": "TE_ESSERR", "CFURH": "TE_CFURH", "CDE Covid TE": "TE_CDE_COVID",
    "CDE ELET": "TE_CDE_ELET", "CDE GD": "TE_CDE_GD", "ENERGIA REVENDA": "TE_ENERGIA",
    "ITAIPU": "TE_TRANSPORTE_ITAIPU", "TUST ITAIPU": "TE_TUST_ITAIPU", "TUST CI": "TE_TUST_CI",
    "SUBSIDIO": "TE Subsidio", "PERDAS RB/C": "TE_Per_RB",
}


def s(v):
    return "" if v is None else str(v).strip()


def num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def le_aba(caminho, aba):
    """Read one PCAT sheet. Header (row 3) drives which column is which --
    never a fixed index -- so the sheet can gain or lose item columns between
    vintages without breaking the read."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[aba]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()

    cabecalho = linhas[2]
    idx = {s(v): i for i, v in enumerate(cabecalho) if s(v)}
    faltando = [a for a in ATRIBUTOS + TOTAIS if a not in idx]
    if faltando:
        raise ValueError(f"{caminho.name} / {aba}: colunas ausentes no cabecalho: {faltando}")

    reservadas = set(ATRIBUTOS) | set(TOTAIS)
    itens = [s(v) for v in cabecalho if s(v) and s(v) not in reservadas]

    registros = []
    for r in linhas[3:]:
        if not s(r[idx["SUBGRUPO"]]):
            continue
        reg = {a: s(r[idx[a]]) for a in ATRIBUTOS}
        for t in TOTAIS:
            reg[t] = num(r[idx[t]])
        for it in itens:
            reg[it] = num(r[idx[it]])
        registros.append(reg)
    return pd.DataFrame(registros), itens


def deriva_desconto_social(tabela):
    """Add the Desconto Social rows to the vintages whose PCAT does not have them."""
    tabela = tabela.copy()
    tabela["ORIGEM"] = "PCAT"
    nativas = tabela[tabela["SUBCLASSE"].isin([DS_FAIXA_01, DS_FAIXA_02])]
    modalidades = sorted(nativas["MODALIDADE"].unique())
    if not modalidades:
        raise ValueError("nenhuma vigencia traz o Desconto Social - sem referencia de modalidades")

    base = tabela[(tabela["SUBGRUPO"] == "B1") & (tabela["SUBCLASSE"] == "Residencial")
                  & tabela["MODALIDADE"].isin(modalidades)]
    anos_nativos = set(nativas["ANO_VIGENCIA"])
    chave = ["ANO_VIGENCIA", "CAMADA"] + ATRIBUTOS
    novas = []
    for ano, g in base.groupby("ANO_VIGENCIA"):
        f2 = g.copy()
        f2["SUBCLASSE"] = DS_FAIXA_02
        f1 = g.copy()
        f1["SUBCLASSE"] = DS_FAIXA_01
        for it in ISENTOS_DESCONTO_SOCIAL:
            total = "Total TE" if it.startswith("TE") else "Total TUSD"
            f1[total] = f1[total] - f1[it]
            f1[it] = 0.0
        derivadas = pd.concat([f1, f2], ignore_index=True)
        derivadas["ORIGEM"] = "DERIVADA_RESIDENCIAL"

        if ano in anos_nativos:
            # check only: how far the rule is from the published rows
            cmp_ = derivadas.merge(nativas, on=chave, suffixes=("_DER", "_PCAT"))
            ta = cmp_[cmp_["CAMADA"] == "TA"]
            for t in TOTAIS:
                d = (ta[f"{t}_DER"] - ta[f"{t}_PCAT"]).abs().max()
                print(f"{ano}: Desconto Social publicado no PCAT - regra derivada x PCAT, "
                      f"{t} (TA), diferenca maxima {d:.2f} R$/MWh")
            continue

        novas.append(derivadas)
        print(f"{ano}: Desconto Social ausente no PCAT - {len(derivadas)} linhas derivadas da "
              f"Residencial ({', '.join(modalidades)})")

    if novas:
        tabela = pd.concat([tabela] + novas, ignore_index=True)
    return tabela


def le_te_suprimento(caminho):
    """{(camada, linha): {item: R$/MWh}} of the supply rows of the TE sheets. Columns located by the
    item label (row 4), never by position; the row total (last column) is kept as a check."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    out = {}
    try:
        for camada, aba in ABAS_TE_SUPRIMENTO.items():
            if aba not in wb.sheetnames:
                raise ValueError(f"{caminho.name}: aba {aba} ausente")
            linhas = list(wb[aba].iter_rows(values_only=True))
            rotulos = [s(v) for v in linhas[3]]
            col = {ITEM_TE[r]: j for j, r in enumerate(rotulos) if r in ITEM_TE}
            faltam = sorted(set(ITEM_TE.values()) - set(col))
            if faltam:
                raise ValueError(f"{caminho.name} / {aba}: itens ausentes no cabecalho: {faltam}")
            total = max(j for j, v in enumerate(linhas[2]) if s(v))   # row total: last labelled column of row 3
            for r in linhas[4:]:
                nome = s(r[1])
                if nome.startswith("Energia suprimento") and "<500" in s(r[4]):
                    valores = {it: num(r[j]) for it, j in col.items()}
                    soma = sum(valores.values())
                    if abs(soma - num(r[total])) > 0.01:
                        raise ValueError(f"{caminho.name} / {aba} / {nome}: itens somam {soma:.4f}, "
                                         f"total da linha {num(r[total]):.4f}")
                    out[(camada, nome)] = valores
    finally:
        wb.close()
    return out


def completa_te_suprimento(tabela):
    """Fill the TE items of the A4 Distribuição MWh keys whose TE is zero, from the supply rows of
    the TE sheets of the same vintage (layers BE, CVA, FIN; TA = BE + CVA + FIN)."""
    tabela = tabela.copy()
    for v in VIGENCIAS:
        alvo = ((tabela["ANO_VIGENCIA"] == v["ano"]) & (tabela["MODALIDADE"] == "Distribuição")
                & (tabela["UNIDADE"] == "MWh"))
        sem_te = tabela.loc[alvo & (tabela["CAMADA"] == "TA") & (tabela["Total TE"].abs() < 1e-9), "ACESSANTE"]
        if sem_te.empty:
            continue
        sup = le_te_suprimento(v["arquivo"])
        for acessante in sorted(sem_te.unique()):
            linha = LINHA_SUPRIMENTO.get(acessante, LINHA_SUPRIMENTO_PADRAO)
            ta = dict.fromkeys(ITEM_TE.values(), 0.0)
            for camada in ("BE", "CVA", "FIN"):
                if (camada, linha) not in sup:
                    raise ValueError(f"{v['ano']}: linha '{linha}' ausente na aba {ABAS_TE_SUPRIMENTO[camada]}")
                vals = sup[(camada, linha)]
                m = alvo & (tabela["CAMADA"] == camada) & (tabela["ACESSANTE"] == acessante)
                for it, x in vals.items():
                    tabela.loc[m, it] = x
                    ta[it] += x
                tabela.loc[m, "Total TE"] = sum(vals.values())
                tabela.loc[m, "ORIGEM"] = "TE_SUPRIMENTO"
            m = alvo & (tabela["CAMADA"] == "TA") & (tabela["ACESSANTE"] == acessante)
            for it, x in ta.items():
                tabela.loc[m, it] = x
            tabela.loc[m, "Total TE"] = sum(ta.values())
            tabela.loc[m, "ORIGEM"] = "TE_SUPRIMENTO"
            print(f"{v['ano']}: TE de suprimento completada para {acessante} ({linha}): "
                  f"TA {sum(ta.values()):.2f} R$/MWh")
    return tabela


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    blocos = []
    itens_por_ano = {}
    pior_fin_geral = 0.0

    for v in VIGENCIAS:
        camadas, itens_por_aba = {}, {}
        for camada, aba in ABAS.items():
            df, itens = le_aba(v["arquivo"], aba)
            camadas[camada] = df
            itens_por_aba[camada] = set(itens)

        n = {c: len(camadas[c]) for c in camadas}
        if len(set(n.values())) != 1:
            raise ValueError(f"{v['ano']}: numero de linhas nao bate entre TA/BE/CVA: {n}")

        # duplicate-key check per layer
        for camada, df in camadas.items():
            dups = df.duplicated(subset=ATRIBUTOS).sum()
            if dups:
                print(f"ATENCAO {v['ano']} / {camada}: {dups} chave(s) duplicada(s) na aba")

        itens_desta_vigencia = sorted(set().union(*itens_por_aba.values()))
        itens_por_ano[v["ano"]] = set(itens_desta_vigencia)
        for camada, df in camadas.items():
            for it in itens_desta_vigencia:
                if it not in df.columns:
                    df[it] = 0.0

        # FIN = TA - BE - CVA, merged on the 8-attribute key (never positional)
        merge_cols = ATRIBUTOS
        m = camadas["TA"].merge(camadas["BE"], on=merge_cols, suffixes=("", "_BE")) \
                          .merge(camadas["CVA"], on=merge_cols, suffixes=("", "_CVA"))
        fin = camadas["TA"][ATRIBUTOS].copy()
        for it in TOTAIS + itens_desta_vigencia:
            fin[it] = m[it] - m[f"{it}_BE"] - m[f"{it}_CVA"]
        camadas["FIN"] = fin

        # validation: TA = BE + CVA + FIN, item a item, for this vintage
        chk = camadas["TA"].merge(camadas["BE"], on=merge_cols, suffixes=("", "_BE")) \
                            .merge(camadas["CVA"], on=merge_cols, suffixes=("", "_CVA")) \
                            .merge(camadas["FIN"], on=merge_cols, suffixes=("", "_FIN"))
        pior = 0.0
        for it in TOTAIS + itens_desta_vigencia:
            d = (chk[it] - (chk[f"{it}_BE"] + chk[f"{it}_CVA"] + chk[f"{it}_FIN"])).abs().max()
            pior = max(pior, d)
        pior_fin_geral = max(pior_fin_geral, pior)
        print(f"{v['ano']}: {len(camadas['TA'])} chaves, {len(itens_desta_vigencia)} itens | "
              f"TA=BE+CVA+FIN, diferenca maxima {pior:.10f}")

        for camada, df in camadas.items():
            df = df.copy()
            df.insert(0, "CAMADA", camada)
            df.insert(0, "VIGENCIA_FIM", v["fim"])
            df.insert(0, "VIGENCIA_INICIO", v["inicio"])
            df.insert(0, "ANO_VIGENCIA", v["ano"])
            blocos.append(df)

    # report item drift between consecutive vintages
    anos = sorted(itens_por_ano)
    for i in range(1, len(anos)):
        prev, cur = anos[i - 1], anos[i]
        novos = itens_por_ano[cur] - itens_por_ano[prev]
        sumidos = itens_por_ano[prev] - itens_por_ano[cur]
        if novos:
            print(f"itens novos em {cur} (nao existiam em {prev}): {sorted(novos)}")
        if sumidos:
            print(f"itens que sumiram em {cur} (existiam em {prev}): {sorted(sumidos)}")
    if pior_fin_geral >= 0.01:
        raise ValueError(f"validacao TA=BE+CVA+FIN falhou: diferenca maxima {pior_fin_geral}")

    todos_itens = sorted(set().union(*itens_por_ano.values()))
    colunas = ["ANO_VIGENCIA", "VIGENCIA_INICIO", "VIGENCIA_FIM", "CAMADA"] + ATRIBUTOS + \
        TOTAIS + todos_itens

    tabela = pd.concat(blocos, ignore_index=True)
    for it in todos_itens:
        if it not in tabela.columns:
            tabela[it] = 0.0
    tabela = tabela[colunas].fillna(0.0)
    tabela = deriva_desconto_social(tabela)
    tabela = completa_te_suprimento(tabela)

    # sanity: vigencias contiguous, no gap or overlap
    ordenadas = sorted(VIGENCIAS, key=lambda v: v["inicio"])
    for a, b in zip(ordenadas, ordenadas[1:]):
        gap = (b["inicio"] - a["fim"]).days
        if gap != 1:
            print(f"ATENCAO: vigencia {a['ano']} termina em {a['fim']} e {b['ano']} comeca em "
                  f"{b['inicio']} -- {'sobreposicao' if gap < 1 else 'buraco'} de {1 - gap} dia(s)")

    caminho_saida = OUTDIR / "tarifa_long.csv"
    tabela.to_csv(caminho_saida, sep=";", decimal=",", index=False, encoding="utf-8-sig")
    print(f"\ngravado: {caminho_saida}  ({len(tabela)} linhas, {len(tabela.columns)} colunas)")
    print(f"vigencias: {[v['ano'] for v in VIGENCIAS]}")
    print(f"itens (uniao de todas as vigencias): {len(todos_itens)}")


if __name__ == "__main__":
    main()
