# -*- coding: utf-8 -*-
"""Convert the reading calendar workbook (ESCALA AT MT BTr e BT, one sheet per
month) into ONE long table, one row per (month x group x lot), with the reading
date and the consumption cycle as real dates.

Input : src/calendario_leitura_<ano>.xlsx  (one file per year, as received)
Output: tbl/calendario/calendario_leitura.csv  (all months loaded so far)

The calendar may arrive incomplete and be revised during the year. Only the
month sheets present in the workbook are loaded; each one replaces the same
(ANO, MES) rows in the table, and the other months already there are kept.
Only the months with prorated tariff (readjustment months and the cycles that
cross the readjustment date) are actually needed downstream.

Groups (GRUPO):
  MT   - Media Tensao, lots 1 to 17
  BT   - Baixa Tensao, lots 1 to 20
  TELE - Telemedidos, lots 40 and 42. The effective reading is the one in the
         section header ("Leitura Real dd/mm", last day of the previous month),
         not the "Data Leitura" column, which is the processing date.
Cash Power is not exported: it has no billing in the BW base.

Cycle: FIM_CICLO = DATA_LEITURA; INICIO_CICLO = DATA_LEITURA - DIAS_CONSUMO + 1.
In a readjustment month the tariff is prorated by the days of each tariff
vintage inside [INICIO_CICLO, FIM_CICLO] (REN ANEEL 1.000, art. 173 par. 2).

Usage: python gera_calendario_leitura.py 2026
"""
import os
import re
import sys
import datetime as dt
from pathlib import Path

import pandas as pd
import openpyxl

BASE = Path(os.environ.get("PXQ_BASE", r"C:\python_projects\claude\p_x_q\SP\BW"))
OUTDIR = BASE / "tbl" / "calendario"
SAIDA = OUTDIR / "calendario_leitura.csv"

MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
LOTES_ESPERADOS = {"MT": set(range(1, 18)), "BT": set(range(1, 21)), "TELE": {40, 42}}
PRIMEIRA_LINHA = 9  # data starts below the column headers

# Column positions in each month sheet (0-based), fixed by the calendar layout
COL_LOTE, COL_DIAS, COL_LEITURA = 0, 3, 4


def grupo_da_secao(texto):
    """Map a section header (column A) to a group code; None if not a header."""
    t = texto.upper()
    if "DIA TENS" in t:
        return "MT"
    if "BAIXA" in t:
        return "BT"
    if "TELEMED" in t:
        return "TELE"
    if "CASH" in t:
        return "CASH"
    return None


def leitura_real(texto, ano, mes):
    """Parse 'Leitura Real dd/mm' into a date; Dec reading of a January sheet is in the previous year."""
    m = re.search(r"(\d{1,2})/(\d{1,2})", str(texto or ""))
    if not m:
        raise ValueError(f"{MESES[mes - 1]}: 'Leitura Real dd/mm' not found in the Telemedidos header")
    d, mm = int(m.group(1)), int(m.group(2))
    return dt.date(ano - 1 if mm > mes else ano, mm, d)


def para_data(v):
    """Excel cell -> date."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    raise ValueError(f"not a date: {v!r}")


def le_aba(ws, ano, mes):
    """Rows of one month sheet: list of dicts."""
    linhas, grupo, real_tele = [], None, None
    for r in ws.iter_rows(min_row=PRIMEIRA_LINHA, values_only=True):
        a = r[COL_LOTE]
        if a is None:
            continue
        s = str(a).strip()
        # lot numbers are stored as text ('1'); handle them before the section headers
        if s.isdigit():
            if grupo in LOTES_ESPERADOS:
                data = real_tele if grupo == "TELE" else para_data(r[COL_LEITURA])
                dias = int(r[COL_DIAS])
                linhas.append({"ANO": ano, "MES": mes, "GRUPO": grupo, "LOTE": int(s),
                               "DATA_LEITURA": data, "DIAS_CONSUMO": dias,
                               "INICIO_CICLO": data - dt.timedelta(days=dias - 1),
                               "FIM_CICLO": data})
            continue
        g = grupo_da_secao(s)
        if g:
            grupo = g
            if g == "TELE":
                real_tele = leitura_real(r[COL_LEITURA], ano, mes)
    return linhas


def valida(df):
    """Fail on structural problems; report cycle gaps/overlaps as warnings."""
    erros = []
    for (mes, grupo), g in df.groupby(["MES", "GRUPO"]):
        falta = LOTES_ESPERADOS[grupo] - set(g["LOTE"])
        if falta:
            erros.append(f"{MESES[mes - 1]} {grupo}: lots missing {sorted(falta)}")
    if df["LOTE"].groupby([df["MES"], df["GRUPO"], df["LOTE"]]).size().gt(1).any():
        erros.append("duplicated (month, group, lot)")
    fora = df[(df["DIAS_CONSUMO"] < 15) | (df["DIAS_CONSUMO"] > 47)]
    if len(fora):
        erros.append(f"implausible DIAS_CONSUMO:\n{fora.to_string(index=False)}")
    if erros:
        raise ValueError("Calendar validation failed:\n" + "\n".join(erros))

    # Continuity: each cycle should start the day after the previous month's cycle ended
    # (checked only when the previous month is also in the table)
    d = df.sort_values(["GRUPO", "LOTE", "ANO", "MES"]).copy()
    d["SEQ"] = d["ANO"] * 12 + d["MES"]
    d["FIM_ANTERIOR"] = d.groupby(["GRUPO", "LOTE"])["FIM_CICLO"].shift()
    d.loc[d.groupby(["GRUPO", "LOTE"])["SEQ"].diff() != 1, "FIM_ANTERIOR"] = None
    gap = d[d["FIM_ANTERIOR"].notna() &
            (pd.to_datetime(d["INICIO_CICLO"]) != pd.to_datetime(d["FIM_ANTERIOR"]) + pd.Timedelta(days=1))]
    if len(gap):
        print(f"WARNING: {len(gap)} cycles do not start the day after the previous reading:")
        print(gap[["ANO", "MES", "GRUPO", "LOTE", "INICIO_CICLO", "FIM_ANTERIOR"]].to_string(index=False))


def main(ano):
    arquivo = BASE / "src" / f"calendario_leitura_{ano}.xlsx"
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    # Load only the month sheets present in the workbook
    presentes = [(i, m) for i, m in enumerate(MESES, 1) if m in wb.sheetnames]
    if not presentes:
        raise ValueError(f"{arquivo.name}: no month sheet found (expected names {MESES})")
    ausentes = [m for m in MESES if m not in wb.sheetnames]
    if ausentes:
        print(f"Month sheets not in the file (kept as they are in the table): {ausentes}")
    novo = pd.DataFrame([x for i, m in presentes for x in le_aba(wb[m], ano, i)])
    novo["ORIGEM"] = arquivo.name

    # Replace only the (ANO, MES) pairs loaded now; keep every other month
    OUTDIR.mkdir(parents=True, exist_ok=True)
    if SAIDA.exists():
        atual = pd.read_csv(SAIDA, sep=";", encoding="utf-8-sig",
                            parse_dates=["DATA_LEITURA", "INICIO_CICLO", "FIM_CICLO"])
        for c in ("DATA_LEITURA", "INICIO_CICLO", "FIM_CICLO"):
            atual[c] = atual[c].dt.date
        carregados = set(zip(novo["ANO"], novo["MES"]))
        manter = ~pd.Series(list(zip(atual["ANO"], atual["MES"]))).isin(carregados).values
        tab = pd.concat([atual[manter], novo], ignore_index=True)
    else:
        tab = novo
    tab = tab.sort_values(["ANO", "MES", "GRUPO", "LOTE"]).reset_index(drop=True)
    valida(tab)
    tab.to_csv(SAIDA, sep=";", index=False, encoding="utf-8-sig", date_format="%Y-%m-%d")

    print(f"{arquivo.name}: months {[m for _, m in presentes]}, {len(novo)} rows -> {SAIDA}")
    print(novo.groupby("GRUPO").agg(lotes=("LOTE", "nunique"), linhas=("LOTE", "size")).to_string())


if __name__ == "__main__":
    if len(sys.argv) != 2 or not sys.argv[1].isdigit():
        sys.exit("Usage: python gera_calendario_leitura.py <ano>")
    main(int(sys.argv[1]))
