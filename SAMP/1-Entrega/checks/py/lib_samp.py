# -*- coding: utf-8 -*-
"""
Shared library of the SAMP checks (projeto_samp.md). The PxQ SAMP calculation is
in SAMP/1-Entrega/PxQ_SAMP.ipynb; the checks read its output.

Functions:
  le_samp(ref)            raw SAMP rows of Enel SP for one competence month
  carrega_pxq_samp(ref)   PxQ SAMP output (SAMP/1-Entrega/pxq/pxq_samp_<ref>.xlsx)
  familias_mr(ref)        Mercado Receita with families (subsidy spread by abatement)
  carrega_pxq_bw(ref)     BW PxQ output (SP/BW/1-Entrega/pxq/pxq_<ref>.xlsx) with the
                          same measures, at group level
  medidas(df, grao)       sum of the comparison measures by a grain
  escreve_comparacao(...) workbook: one sheet per measure/family, rows OPCAO x SUBGRUPO,
                          columns side A | side B | Delta | Delta % | share of delta

Rules of the SAMP PxQ: PxQ_SAMP.ipynb and projeto_samp.md section 4.1.
Nothing under SP/BW/ is written; tariffs, MR and BW PxQ outputs are only read.
"""
import datetime as dt
import os
import sys
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(os.environ.get("PXQ_ROOT", r"C:\python_projects\claude\p_x_q"))
sys.path.insert(0, str(ROOT / "SP" / "BW" / "3-Checks" / "py"))
from le_mr import carrega_mr as _carrega_mr  # noqa: E402
CNPJ_SP = "61695227000193"
NSA = "Não se aplica"
CAMADAS = ["TA", "BE", "CVA", "FIN"]
FAMILIAS = ["FIO A", "FIO B", "ENCARGOS", "PERDAS", "ENERGIA", "RI", "OUTROS"]
MEDIDAS = ["ENERGIA_TUSD_MWH", "ENERGIA_TE_MWH", "DEMANDA_KW", "MONETARIO_TOTAL"] + FAMILIAS
ORDEM_OPCAO = ["CATIVO", "DISTRIBUIÇÃO", "GERAÇÃO", "LIVRE", "SUPRIMENTO"]
ORDEM_SUBGRUPO = ["A2", "A3a", "A4", "AS", "B1", "B2", "B3", "B4"]

# component -> family (same map as the BW notebook, Passo 6)
FAMILIA = {
    "FIO A": ["TUSD_RB", "TUSD_FR", "TUSD_CCT", "TUSD_CCD", "TUSD_CUSD", "TUSDG_T",
              "TE_TRANSPORTE_ITAIPU", "TE_TUST_ITAIPU", "TE_TUST_CI"],
    "FIO B": ["TUSD_FioB"],
    "ENCARGOS": ["TUSD_CDE_COVID", "TUSD_TFSEE", "TUSD_PeD", "TUSD_ONS", "TUSD_CCC", "TUSD_CDE",
                 "TUSD_PROINFA", "Liminar1", "TUSDG_ONS", "TE_CDE_COVID", "TE_CDE_ELET",
                 "TE_CDE_GD", "TE_PeD", "TE_ESSERR", "TE_CFURH"],
    "PERDAS": ["TUSD_PT", "TUSD_Per_RB_D", "TUSD_PNT", "TE_Per_RB"],
    "ENERGIA": ["TE_ENERGIA", "TE_ANGRA"],
    "RI": ["TUSD_RI"],
    "OUTROS": ["TUSD Subsidio", "TUSD BENEFICIO_L14299", "TUSD Outros",
               "TE Subsidio", "TE BENEFICIO_L14299"],
}

# SAMP accessing agent -> PCAT ACESSANTE (Tecipar has no row of its own: generic key)
ACESSANTE_PCAT = {
    "ELEKTRO REDES S.A.": "Neoenergia Elektro",
    "COMPANHIA PIRATININGA DE FORCA E LUZ": "CPFL Piratininga",
    "COOPERATIVA DE ELETRIFICAÇÃO DA REGIÃO DE ITAPECERICA DA SERRA": "Ceris",
    "Rasgão": "PCH RASGAO",
    "Nova Piratininga (Antiga Fernando Gasparian)": "UTE FERNANDO GASPARIAN (N.PIRAT)",
    "São João Biogás": "UTE SAO JOAO BIOGAS",
    "Pirapora": "PCH PIRAPORA",
}
GRUPO_SAMP = ["TIPO_MERCADO", "OPCAO", "SUBGRUPO", "MODALIDADE", "CLASSE_SAMP",
              "SUBCLASSE_SAMP", "DETALHE_SAMP", "POSTO", "AGENTE"]


# ====================================================================== SAMP
def le_samp(ref):
    """SAMP rows of Enel SP for one competence month (CNPJ compared as 14-digit text)."""
    samp = pd.read_parquet(ROOT / "SAMP" / "data" / f"samp-{ref[:4]}.parquet")
    samp = samp[samp["NumCNPJAgenteDistribuidora"].astype(str).str.zfill(14) == CNPJ_SP]
    samp = samp[pd.to_datetime(samp["DatCompetencia"]).dt.strftime("%Y-%m") == ref].copy()
    assert len(samp), f"SAMP sem linhas para {ref}"
    samp = samp.rename(columns={
        "NomTipoMercado": "TIPO_MERCADO", "DscOpcaoEnergia": "OPCAO",
        "DscSubGrupoTarifario": "SUBGRUPO", "DscModalidadeTarifaria": "MODALIDADE",
        "DscClasseConsumoMercado": "CLASSE_SAMP", "DscSubClasseConsumidor": "SUBCLASSE_SAMP",
        "DscDetalheConsumidor": "DETALHE_SAMP", "DscPostoTarifario": "POSTO",
        "NomAgenteAcessante": "AGENTE"})
    for c in GRUPO_SAMP:
        samp[c] = samp[c].fillna(NSA).astype(str).str.strip()
    return samp


# ====================================================================== PxQ SAMP (notebook output)
# PxQ_SAMP.ipynb labels -> short names used by the checks
_ROTULO_NB = {
    "Source": "FONTE", "Tipo de mercado": "TIPO_MERCADO", "Opção": "OPCAO", "Subgrupo": "SUBGRUPO",
    "Modalidade": "MODALIDADE", "Classe (SAMP)": "CLASSE_SAMP", "Subclasse (SAMP)": "SUBCLASSE_SAMP",
    "Detalhe (SAMP)": "DETALHE_SAMP", "Posto": "POSTO", "Agente acessante (SAMP)": "AGENTE",
    "Classe (tarifa)": "CLASSE", "Subclasse (tarifa)": "SUBCLASSE", "Detalhe (tarifa)": "DETALHE",
    "Acessante (tarifa)": "ACESSANTE", "Chave energia": "CHAVE_E",
    "Regra da chave de energia": "REGRA_E", "Chave demanda": "CHAVE_D",
    "Regra da chave de demanda": "REGRA_D", "Energia TUSD (MWh)": "E_TUSD_MWH",
    "Energia TE (MWh)": "E_TE_MWH", "Demanda (kW)": "DEMANDA_KW",
    "Desconto demanda % (SAMP)": "DESC_DEM_PCT", "Desconto energia TUSD % (SAMP)": "DESC_E_PCT",
    "RECEITA SAMP": "RECEITA_BRL", "VALOR CALCULADO (bloco 1)": "BLOCO1_TA",
    "Tarifável?": "TARIFAVEL", "Regime de ajuste": "REGIME", "% de desconto no fio": "DESC_FIO",
    "Ajuste 1ª passada (R$)": "AJUSTE_1A_BRL", "Ajuste 2ª passada (R$)": "AJUSTE_2A_BRL",
    "Resíduo não ajustável (R$)": "RESIDUO_NAO_AJUSTAVEL",
}


def carrega_pxq_samp(ref):
    """PxQ SAMP output of PxQ_SAMP.ipynb (SAMP/1-Entrega/pxq/pxq_samp_<ref>.xlsx, tab Familias),
    one row per SAMP group, with B1_<family> / B2_<family> (layer TA) and the comparison
    measures (bloco 2). The calculation itself lives only in the notebook."""
    arq = ROOT / "SAMP" / "1-Entrega" / "pxq" / f"pxq_samp_{ref}.xlsx"
    if not arq.exists():
        raise FileNotFoundError(f"falta {arq.name} - rodar SAMP/1-Entrega/PxQ_SAMP.ipynb para {ref}")
    fam = pd.read_excel(arq, sheet_name="Familias")
    base = fam.rename(columns=_ROTULO_NB)
    for c in GRUPO_SAMP + ["CLASSE", "SUBCLASSE", "DETALHE", "ACESSANTE", "REGRA_E", "REGRA_D"]:
        base[c] = base[c].fillna("" if c.startswith("REGRA") else NSA)
    for f in FAMILIAS:
        base[f"B1_{f}"] = fam[f"[1] TA · {f}"].fillna(0.0)
        base[f"B2_{f}"] = fam[f"[2] TA · {f}"].fillna(0.0)
    base["BLOCO2_TA"] = base[[f"B2_{f}" for f in FAMILIAS]].sum(axis=1)
    base["ENERGIA_TUSD_MWH"] = base["E_TUSD_MWH"]
    base["ENERGIA_TE_MWH"] = base["E_TE_MWH"]
    base["MONETARIO_TOTAL"] = base["BLOCO2_TA"]
    for f in FAMILIAS:
        base[f] = base[f"B2_{f}"]
    return base


# ====================================================================== MR
FAMILIA_MR = {
    "FIO A": ["TUST_RB", "CVA_TUST_RB", "finCT_TUST_RB", "TUST_FR", "CVA_TUST_FR", "finCT_TUST_FR",
              "Conexao_T", "FinCT_Conexao_T", "Conexao_D", "FinCT_Conexao_D", "CUSD", "FinCT_CUSD",
              "TUSDG_T", "CVA_TUSDG_T", "FINCT_TUSDG_T", "ITAIPU", "CVA_TranItaipu", "FinCT_Itaipu",
              "TUST_ITAIPU", "CVA_TUST_ITAIPU", "FINCT_TUST_ITAIPU", "TUST_CI", "CVA_TUST_CI",
              "FINCT_TUST_CI"],
    "FIO B": ["DISTRIBUICAO", "FinCT_ParcelaB"],
    "ENCARGOS": ["covidTusdEco", "CVA_CovidTusd", "FinCT_CDECOVIDTUSD", "TFSEE", "finCT_TFSEE",
                 "PeD", "FinCT_Ped", "ONS", "FinCT_ONS", "CCC", "CVA_CCC", "FinCT_CCC",
                 "CDE", "CVA_CDE", "FinCT_CDE", "PROINFA", "CVA_PROINFA", "finCT_PROINFA",
                 "CVA_liminar", "FINCT_Liminar", "TUSDG_ONS", "CVA_TUSDG_ONS", "FINCT_TUSDG_ONS",
                 "covidTeEco", "CVA_CovidTe_ou_CVA_CDE_E", "FinCT_CDECOVIDTE_ou_FinCT_CDE_E",
                 "CDE_Elet_Eco", "CVA_CDE_Elet", "FinCT_CDE_Elet", "CDE_GD_Eco", "CVA_CDE_GD",
                 "FinCT_CDE_GD", "ESSERR", "CVA_ESSERR", "FinCT_ESSERR", "CFURH", "CVA_CFURH",
                 "FinCT_CFURH", "RGR", "FINCT_RGR", "CDE_E", "CVA_CDE_E", "FinCT_CDE_E"],
    "PERDAS": ["perdastec", "CVA_perdastec", "FinCT_perdastec", "perdasrb_d", "CVA_perdasrb_d",
               "FINCT_perdasrb_d", "perdasnt", "CVA_perdasnt", "FinCT_perdasnt",
               "perdasrb_c", "CVA_perdasrb_c", "FINCT_perdasrb_c"],
    "ENERGIA": ["EnergiaRevenda", "CVA_EnergiaRevenda", "FinCT_EnergiaRevenda",
                "ANGRA_Eco", "CVA_ANGRA", "FinCT_ANGRA"],
    "RI": ["RI", "FINCT_RI"],
    "OUTROS": ["OUTROS", "FinCT_Subsidio", "Subs_BENEFICIO_L14299_TUSD", "Subs_BENEFICIO_L14299_TE"],
}
PARCELA_TE_MR = {
    "EnergiaRevenda", "CVA_EnergiaRevenda", "FinCT_EnergiaRevenda", "ANGRA_Eco", "CVA_ANGRA",
    "FinCT_ANGRA", "ESSERR", "CVA_ESSERR", "FinCT_ESSERR", "CFURH", "CVA_CFURH", "FinCT_CFURH",
    "covidTeEco", "CVA_CovidTe_ou_CVA_CDE_E", "FinCT_CDECOVIDTE_ou_FinCT_CDE_E",
    "CDE_Elet_Eco", "CVA_CDE_Elet", "FinCT_CDE_Elet", "CDE_GD_Eco", "CVA_CDE_GD", "FinCT_CDE_GD",
    "CDE_E", "CVA_CDE_E", "FinCT_CDE_E", "perdasrb_c", "CVA_perdasrb_c", "FINCT_perdasrb_c",
    "ITAIPU", "CVA_TranItaipu", "FinCT_Itaipu", "TUST_ITAIPU", "CVA_TUST_ITAIPU",
    "FINCT_TUST_ITAIPU", "TUST_CI", "CVA_TUST_CI", "FINCT_TUST_CI", "Subs_BENEFICIO_L14299_TE",
}


def familias_mr(ref, incluir_erc=True):
    """MR rows with OPCAO, SUBGRUPO, the comparison measures and the families
    (subsidy spread by the MR quantity abatement, as in check_mr_bloco2.py)."""
    # same MR loader as the BW checks: single-month file or tariff-period file,
    # PROJETADO months refused (SP/BW/3-Checks/py/le_mr.py)
    mr, _ = _carrega_mr(ROOT / "SP" / "BW", ref)
    mr = mr[mr["Subgrupo"].notna()].copy()
    if not incluir_erc:
        mr = mr[mr["OPÇÃO"] != "ERC"].copy()
    mr["OPCAO"] = mr["OPÇÃO"].replace({"LIVRE APE": "LIVRE", "ERC": "LIVRE"})
    mr["SUBGRUPO"] = mr["Subgrupo"]
    for c in ["D", "Daj", "TUSD_E", "TUSD_Eaj", "TE_E", "TE_Eaj", "FinTotAj",
              "Total com Tar de aplicação", "SubsidioTotal", "Subs_CarFI"]:
        mr[c] = mr[c].fillna(0.0)
    abate_d = (mr["D"] - mr["Daj"]).abs() > 1e-6
    abate_e = (mr["TUSD_E"] - mr["TUSD_Eaj"]).abs() > 1e-6
    assert not (abate_d & (mr["TUSD_E"] != 0) | abate_e & (mr["D"] != 0)).any()
    mr["FATOR_TUSD"] = np.nan_to_num(np.where(
        mr["D"] != 0, mr["Daj"] / mr["D"].replace(0, np.nan),
        np.where(mr["TUSD_E"] != 0, mr["TUSD_Eaj"] / mr["TUSD_E"].replace(0, np.nan), 1.0)), nan=1.0)
    mr["FATOR_TE"] = np.nan_to_num(np.where(
        mr["TE_E"] != 0, mr["TE_Eaj"] / mr["TE_E"].replace(0, np.nan), 1.0), nan=1.0)
    bruto = mr[[c for l in FAMILIA_MR.values() for c in l]].fillna(0.0).sum(axis=1)
    for fam, lista in FAMILIA_MR.items():
        te = [c for c in lista if c in PARCELA_TE_MR]
        tu = [c for c in lista if c not in PARCELA_TE_MR]
        mr[fam] = (mr[tu].fillna(0.0).sum(axis=1) * mr["FATOR_TUSD"]
                   + mr[te].fillna(0.0).sum(axis=1) * mr["FATOR_TE"])
    mr["NAO_ALOCADO"] = mr["Total com Tar de aplicação"] - bruto
    mr["SUBSIDIO_NAO_DISTRIBUIDO"] = mr["FinTotAj"] - mr[FAMILIAS].sum(axis=1) - mr["NAO_ALOCADO"]
    mr["ENERGIA_TUSD_MWH"] = mr["TUSD_E"]
    mr["ENERGIA_TE_MWH"] = mr["TE_E"]
    mr["DEMANDA_KW"] = mr["D"]
    mr["MONETARIO_TOTAL"] = mr["FinTotAj"]
    return mr


# ====================================================================== BW PxQ
OPCAO_BW = {"Cliente Cativo": "CATIVO", "Cliente Livre": "LIVRE", "Cliente Híbrido": "LIVRE",
            "GERAÇÃO": "GERAÇÃO", "USO DISTRIBUIÇÃO": "DISTRIBUIÇÃO", "SUPRIMENTO": "SUPRIMENTO"}
_PADRAO = re.compile(r"^(A2|A3a|A4|AS|B1|B2|B3|B4)(Azul|Verde|Branca|Convencional|Geração|Distribuição)")


def carrega_pxq_bw(ref):
    """BW PxQ output, tariff-eligible groups, with OPCAO / SUBGRUPO / MODALIDADE,
    the comparison measures (bloco 2) and the source flag. Same reclassifications
    as check_mr_bloco2.py: Hibrido -> LIVRE; SD generators on a load tariff
    (EMAE pumping stations) -> LIVRE. Energy: captive TUSD = TE = booked quantity;
    free market all on TUSD, TE = 0; SD rows keep their own sides."""
    arq = ROOT / "SP" / "BW" / "1-Entrega" / "pxq" / f"pxq_{ref}.xlsx"
    comp = pd.read_excel(arq, sheet_name="Componentes")
    fam = pd.read_excel(arq, sheet_name="Familias")
    assert len(comp) == len(fam) and (comp["Chave"].fillna("") == fam["Chave"].fillna("")).all()
    ok = (comp["Tarifável?"] == "SIM") & comp["Chave"].notna()
    df = comp.loc[ok].copy()
    for f in FAMILIAS:
        df[f] = fam.loc[ok, f"[2] TA · {f}"].fillna(0.0).values
    df["FLAG"] = df["Flag Cli.Liv./ Cat."]
    df["OPCAO"] = df["FLAG"].map(OPCAO_BW)
    assert df["OPCAO"].notna().all(), df.loc[df["OPCAO"].isna(), "FLAG"].unique()
    m = df["Chave"].str.extract(_PADRAO)
    df["SUBGRUPO"], df["MODALIDADE"] = m[0], m[1]
    assert df["SUBGRUPO"].notna().all()
    carga = (df["OPCAO"] == "GERAÇÃO") & (df["MODALIDADE"] != "Geração")
    df.loc[carga, "OPCAO"] = "LIVRE"
    df["GRANDEZA"] = np.where(df["Chave"].str.endswith("kW"), "DEMANDA", "CONSUMO")
    tusd = df["Consumo TUSD (kWh)"].fillna(0.0) / 1000.0
    te = df["Consumo TE (kWh)"].fillna(0.0) / 1000.0
    maior = np.maximum(tusd, te)
    cativo = df["FLAG"] == "Cliente Cativo"
    livre = df["FLAG"].isin(["Cliente Livre", "Cliente Híbrido"])
    df["ENERGIA_TUSD_MWH"] = np.where(cativo | livre, maior, tusd)
    df["ENERGIA_TE_MWH"] = np.where(cativo, maior, np.where(livre, 0.0, te))
    df["DEMANDA_KW"] = df["Demanda (kW)"].fillna(0.0)
    df["MONETARIO_TOTAL"] = df["[2] Total TA"].fillna(0.0)
    df["RECEITA_BRL"] = df["RECEITA DE FORNECIMENTO"].fillna(0.0)
    df["FONTE"] = df["Source"]
    return df


# ====================================================================== comparison
def medidas(df, grao=("OPCAO", "SUBGRUPO")):
    return df.groupby(list(grao))[MEDIDAS].sum()


def alinha(a, b):
    idx = a.index.union(b.index)
    chaves = sorted(idx, key=lambda k: (ORDEM_OPCAO.index(k[0]) if k[0] in ORDEM_OPCAO else 99,
                                        ORDEM_SUBGRUPO.index(k[1]) if k[1] in ORDEM_SUBGRUPO else 99))
    return a.reindex(chaves, fill_value=0.0), b.reindex(chaves, fill_value=0.0), chaves


_AZUL = PatternFill("solid", start_color="1F3864")
_CINZA = PatternFill("solid", start_color="D9E1F2")
_CINZA2 = PatternFill("solid", start_color="B4C6E7")
_FINA = Side(style="thin", color="BFBFBF")
_BORDA = Border(left=_FINA, right=_FINA, top=_FINA, bottom=_FINA)
_FN = "Arial"
_NUM = '#,##0;[Red]-#,##0;"-"'
_PCT = '0.0%;[Red]-0.0%;"-"'
ABAS = [
    ("1_Energia_TUSD_MWh", "ENERGIA_TUSD_MWH", "Energia TUSD (MWh)", "MWh"),
    ("2_Energia_TE_MWh", "ENERGIA_TE_MWH", "Energia TE (MWh)", "MWh"),
    ("3_Demanda_kW", "DEMANDA_KW", "Demanda (kW)", "kW"),
    ("4_Monetario_Total_BRL", "MONETARIO_TOTAL", "Monetário total (R$)", "R$"),
] + [(f"{5 + i}_{f.replace(' ', '_')}", f, f"Família {f} (R$)", "R$") for i, f in enumerate(FAMILIAS)]


def _cabecalho(ws, linha, titulos):
    for j, tt in enumerate(titulos, 1):
        c = ws.cell(row=linha, column=j, value=tt)
        c.font = Font(name=_FN, bold=True, color="FFFFFF")
        c.fill = _AZUL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDA


def _formata(ws, r, ncol, destaque=None, num_cols=(3, 4, 5), pct_cols=(6, 7)):
    for j in range(1, ncol + 1):
        c = ws.cell(row=r, column=j)
        c.font = Font(name=_FN, bold=destaque is not None or c.font.bold, size=10)
        c.border = _BORDA
        if destaque:
            c.fill = destaque
        c.number_format = _NUM if j in num_cols else (_PCT if j in pct_cols else "General")


def escreve_comparacao(saida, titulo, nome_a, nome_b, med_a, med_b, chaves,
                       conciliacao=(), memoria=(), detalhe=None, nome_detalhe="Base"):
    """Workbook: 0_Resumo + one sheet per measure/family (rows OPCAO x SUBGRUPO,
    subtotal per OPCAO, TOTAL; Delta, Delta % and share are Excel formulas).
    conciliacao: rows (label, value A or formula, value B or formula) appended to the
    summary after the family block; memoria: rows (label, value A, value B or None)."""
    wb = Workbook()
    wb.remove(wb.active)
    ws0 = wb.create_sheet("0_Resumo")
    totais = {}
    for nome, medida, tt, unidade in ABAS:
        ws = wb.create_sheet(nome)
        ws.cell(row=1, column=1, value=f"{tt} - {titulo}").font = Font(name=_FN, bold=True, size=12)
        _cabecalho(ws, 3, ["OPCAO", "SUBGRUPO", nome_a, nome_b, "Delta", "Delta %",
                           "Participação no delta total"])
        linha, subtotais = 4, []
        for opcao in [o for o in ORDEM_OPCAO if o in {k[0] for k in chaves}]:
            ini = linha
            for k in [c for c in chaves if c[0] == opcao]:
                ws.cell(row=linha, column=1, value=k[0])
                ws.cell(row=linha, column=2, value=k[1])
                ws.cell(row=linha, column=3, value=float(med_a.loc[k, medida]))
                ws.cell(row=linha, column=4, value=float(med_b.loc[k, medida]))
                linha += 1
            ws.cell(row=linha, column=1, value=f"Subtotal {opcao}")
            for col in (3, 4):
                L = get_column_letter(col)
                ws.cell(row=linha, column=col, value=f"=SUM({L}{ini}:{L}{linha - 1})")
            subtotais.append(linha)
            linha += 1
        tl = linha
        ws.cell(row=tl, column=1, value="TOTAL")
        for col in (3, 4):
            L = get_column_letter(col)
            ws.cell(row=tl, column=col, value="=" + "+".join(f"{L}{s}" for s in subtotais))
        for r in range(4, tl + 1):
            ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
            ws.cell(row=r, column=6, value=f'=IF(D{r}=0,"",E{r}/D{r})')
            ws.cell(row=r, column=7, value=f'=IF($E${tl}=0,"",E{r}/$E${tl})')
            _formata(ws, r, 7, _CINZA2 if r == tl else (_CINZA if r in subtotais else None))
        for j, w in enumerate([16, 11, 18, 18, 16, 11, 14], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "C4"
        totais[medida] = (nome, tl, tt, unidade)

    ws0.cell(row=1, column=1, value=f"Resumo - {titulo}").font = Font(name=_FN, bold=True, size=12)
    _cabecalho(ws0, 3, ["Medida", "Unidade", nome_a, nome_b, "Delta", "Delta %"])
    r, lin_fam = 4, []
    for medida, (nome, tl, tt, u) in totais.items():
        ws0.cell(row=r, column=1, value=tt)
        ws0.cell(row=r, column=2, value=u)
        ws0.cell(row=r, column=3, value=f"='{nome}'!C{tl}")
        ws0.cell(row=r, column=4, value=f"='{nome}'!D{tl}")
        ws0.cell(row=r, column=5, value=f"=C{r}-D{r}")
        ws0.cell(row=r, column=6, value=f'=IF(D{r}=0,"",E{r}/D{r})')
        _formata(ws0, r, 6, pct_cols=(6,))
        if medida in FAMILIAS:
            lin_fam.append(r)
        r += 1
    lin_mon = 4 + list(totais).index("MONETARIO_TOTAL")
    if conciliacao:
        r += 1
        ws0.cell(row=r, column=1, value="Conciliação das famílias com o monetário total").font = \
            Font(name=_FN, bold=True)
        r += 1
        ini = r
        for rot, va, vb in [("Soma das famílias", f"=SUM(C{lin_fam[0]}:C{lin_fam[-1]})",
                             f"=SUM(D{lin_fam[0]}:D{lin_fam[-1]})")] + list(conciliacao):
            ws0.cell(row=r, column=1, value=rot)
            ws0.cell(row=r, column=2, value="R$")
            ws0.cell(row=r, column=3, value=va)
            ws0.cell(row=r, column=4, value=vb)
            ws0.cell(row=r, column=5, value=f"=C{r}-D{r}")
            _formata(ws0, r, 5)
            r += 1
        ws0.cell(row=r, column=1, value="Soma - monetário total (conferência, deve ser zero)")
        ws0.cell(row=r, column=2, value="R$")
        ws0.cell(row=r, column=3, value=f"=SUM(C{ini}:C{r - 1})-C{lin_mon}")
        ws0.cell(row=r, column=4, value=f"=SUM(D{ini}:D{r - 1})-D{lin_mon}")
        _formata(ws0, r, 4)
        r += 1
    if memoria:
        r += 1
        ws0.cell(row=r, column=1, value="Memória (valores, não fórmulas)").font = Font(name=_FN, bold=True)
        r += 1
        for rot, va, vb in memoria:
            ws0.cell(row=r, column=1, value=rot)
            ws0.cell(row=r, column=2, value="R$")
            ws0.cell(row=r, column=3, value=float(va))
            if vb is not None:
                ws0.cell(row=r, column=4, value=float(vb))
                ws0.cell(row=r, column=5, value=f"=C{r}-D{r}")
                ws0.cell(row=r, column=6, value=f'=IF(D{r}=0,"",E{r}/D{r})')
            _formata(ws0, r, 6 if vb is not None else 3, pct_cols=(6,))
            r += 1
    for j, w in enumerate([72, 9, 18, 18, 16, 11], 1):
        ws0.column_dimensions[get_column_letter(j)].width = w

    if detalhe is not None:
        wsd = wb.create_sheet(nome_detalhe)
        _cabecalho(wsd, 1, list(detalhe.columns))
        for i, row in enumerate(detalhe.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                wsd.cell(row=i, column=j, value=None if (isinstance(v, float) and np.isnan(v)) else v)
        wsd.freeze_panes = "A2"
    wb.calculation.fullCalcOnLoad = True
    wb.save(saida)
